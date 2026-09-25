#!/usr/bin/env python3
"""Vendor-authored, deterministic CLI comparison, with independent file oracles.

No golden files and no reference tool used as a judge. Dependencies are stdlib
and the existing stdlib synthetic builders. Downloads require --download.
"""
from __future__ import annotations

import argparse
from collections import Counter
import csv
import copy
from datetime import datetime, timezone
import hashlib
import io
import json
import math
import os
from pathlib import Path
import posixpath
import re
import shutil
import signal
import statistics
import struct
import subprocess
import sys
import tarfile
import tempfile
import time
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
import zlib

REPO = Path(__file__).resolve().parents[2]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))
from tools.cli.evaluate_sample import (
    TABLE_PARTS, WORKBOOK_PARTS, DECK_PARTS, stored_zip, stored_png,
    classic_pdf, pdf_current_page_texts, pdf_current_objects, pdf_object, pdf_reference, resolved_text,
)
from tools.cli.artifact_identity import launcher_artifact
from tools.corpus import package_semantics

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'
NS = dict(w=W, s=S, p=P, a=A, r=R)
APP = REPO / 'desktop-app/build/compose/binaries/main/app/Universal File Opener/bin/ufo'
TIMEOUT = 120
MAX_CAPTURE = 8 * 1024 * 1024
MAX_PACKAGE = 16 * 1024 * 1024
AUTHOR = 'Comparison reviewer'
OLD = 'Revenue up 10%'
NEW = 'Revenue up 12%'
COMMENT = 'Please verify the revenue figure.'
INSERTED = 'Added review paragraph.'
PDF_TEXTS = ['Appendix page one', 'Appendix page two', 'Appendix page three']
# Pinned before measurement. These are input intentions, never recorded tool output.
# DOCX indices here count direct body paragraphs from zero; cell is table 0,row 1,col 0.
TASKS = (
    dict(id='docx_read', family='docx', operation='read', paragraph=1),
    dict(id='docx_replace', family='docx', operation='replace', paragraph=1, before=OLD, after=NEW, preserve_runs=True),
    dict(id='docx_insert', family='docx', operation='insert', paragraph=1, text=INSERTED),
    dict(id='docx_delete', family='docx', operation='delete', paragraph=2),
    dict(id='docx_comment', family='docx', operation='comment', paragraph=1, text=COMMENT, author=AUTHOR),
    dict(id='docx_track', family='docx', operation='track', paragraph=1, before=OLD, after=NEW, author=AUTHOR),
    dict(id='docx_cell', family='docx', operation='cell', table=0, row=1, column=0, before='Acme', after='Nova'),
    dict(id='docx_outline', family='docx', operation='outline'),
    dict(id='xlsx_read', family='xlsx', operation='read', sheet='Summary', range='A1:D1'),
    dict(id='xlsx_number', family='xlsx', operation='number', sheet='Summary', cell='B1', before='10', after='12.5'),
    dict(id='xlsx_boolean', family='xlsx', operation='boolean', sheet='Summary', cell='C1', before='TRUE', after='false'),
    dict(id='xlsx_string', family='xlsx', operation='string', sheet='Summary', cell='E1', before=None, after='0012'),
    dict(id='xlsx_formula', family='xlsx', operation='formula', sheet='Summary', cell='E1', before=None, after='B1+1'),
    dict(id='xlsx_append', family='xlsx', operation='append', sheet='Summary', values={'A':'string:Nova','B':'number:12.5','C':'boolean:false'}),
    dict(id='xlsx_rename', family='xlsx', operation='rename', sheet='Summary', after='Reviewed'),
    dict(id='pptx_read', family='pptx', operation='read', slide=1),
    dict(id='pptx_replace', family='pptx', operation='replace', slide=1, shape_id=3, before=OLD, after=NEW, preserve_runs=True),
    dict(id='pptx_notes', family='pptx', operation='notes', slide=1, before='Mention the new client', after='Thank the new client'),
    dict(id='pptx_duplicate', family='pptx', operation='duplicate', slide=1, after=1),
    dict(id='pptx_delete', family='pptx', operation='delete', slide=2),
    dict(id='pdf_read', family='pdf', operation='read', pages=[1,2,3]),
    dict(id='pdf_keep', family='pdf', operation='keep', pages=[3,1]),
    dict(id='refuse_preimage', family='docx', operation='preimage', paragraph=1, before='This preimage is stale', after=NEW),
    dict(id='refuse_empty_cell', family='docx', operation='empty_cell', table=0, row=1, column=0, before='Acme'),
)
MANIFEST_VERSION = 1
# OfficeCLI v1.0.150 (faceb42654004f1fa5c40fb0ce641c42b7dc5a2beb270f25971ea6265b7dc227) is the
# pin the September 15 to 17 evidence names; v1.0.152 replaced it on September 24.
# GenOffice ships its command line inside the desktop package: the pin is the Debian package, the
# tool is the `genoffice` launcher inside it, run in place on the package's own Electron runtime.
PINS = {
    'officecli': dict(tag='v1.0.152', repository='iOfficeAI/OfficeCLI', asset='officecli-linux-x64',
                      sha256='e54d3c1d248372365f0634aac56d6f1918bd04d6e71afc792ad50e075f56cfe9', license='Apache-2.0'),
    'docx-cli': dict(tag='v0.25.0', repository='kklimuk/docx-cli', asset='docx-linux-x64',
                    sha256='e59d32f2a1ffd696bbb816015bea1f437cba4f3864e0e62f6b83df9acc55bfe6', license='MIT'),
    'genoffice': dict(tag='v0.10.1038', repository='genspark-ai/genoffice', asset='genoffice_0.10.1038_amd64.deb',
                      sha256='feee2e0a291cddd9df838f396724a156af57aec65161832c94a08bfa882cc8ef', license='Apache-2.0',
                      package='deb', entry='opt/GenOffice/resources/cli/genoffice',
                      verify=('opt/GenOffice/resources/cli/genoffice', 'opt/GenOffice/resources/cli/genoffice.cjs',
                              'opt/GenOffice/genoffice')),
}
MAX_DOWNLOAD = 160 * 1024 * 1024
for _pin in PINS.values():
    _pin['url'] = f"https://github.com/{_pin['repository']}/releases/download/{_pin['tag']}/{_pin['asset']}"


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def json_bytes(value) -> bytes:
    return (json.dumps(value, sort_keys=True, indent=2, ensure_ascii=False) + '\n').encode()


UNICODE_PATH_EXTRA_ID = 0x7075
UTF8_NAME_FLAG = 0x800


def declared_member_name(info: zipfile.ZipInfo) -> str | None:
    """The name this member's Info-ZIP Unicode path extra field declares, or None.

    A reader that honours the field substitutes that name for the one the ZIP stores, so it is a
    second name the member answers to. The field counts only when it parses as version 1 and its
    name CRC-32 covers the bytes the member really stores for its name, which is the rule the
    extra field carries for exactly this reason. It is read here rather than taken from
    `ZipInfo.filename` so this runner judges a package the same way on every Python: 3.12
    substitutes the declared name and earlier versions do not.
    """
    stored = info.orig_filename.encode('utf-8' if info.flag_bits & UTF8_NAME_FLAG else 'cp437', 'replace')
    extra, at = info.extra, 0
    while at + 4 <= len(extra):
        tag, size = struct.unpack('<HH', extra[at:at + 4])
        body = extra[at + 4:at + 4 + size]
        at += 4 + size
        if tag != UNICODE_PATH_EXTRA_ID or len(body) < 5:
            continue
        version, name_crc = struct.unpack('<BL', body[:5])
        if version != 1 or name_crc != zlib.crc32(stored):
            continue
        try:
            declared = body[5:].decode('utf-8')
        except UnicodeDecodeError:
            continue
        if declared:
            return declared
    return None


def package(data: bytes, maximum: int = MAX_PACKAGE) -> dict[str, bytes]:
    """Every member of the package, keyed by the name the ZIP STORES for it.

    That stored name (`orig_filename`) is the name OPC resolves a part by, the name a receipt's
    write set is written in, and the name every tool under test addresses. A member may also
    answer to a name an Info-ZIP Unicode path extra field declares, which a reader may substitute
    (Python does from 3.12 on); in a package whose entries declare each other's names, that
    points one name at a different member, so a judgement taken from it would compare one member
    of the source against another of the output. Both readings still have to name one member
    each, because a package where either repeats a name cannot be opened by every reader.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        # Directory entries are ZIP bookkeeping, not OPC document parts.
        infos = [i for i in z.infolist() if not i.is_dir()]
        for reading in (lambda i: i.orig_filename,
                        lambda i: declared_member_name(i) or i.orig_filename):
            require(len({reading(i) for i in infos}) == len(infos), 'duplicate ZIP members')
        require(sum(i.file_size for i in infos) <= maximum, 'package size bound exceeded')
        return {i.orig_filename: z.read(i) for i in infos}


def fixtures() -> dict[str, bytes]:
    """Four synthetic lineages, with fixed ZIP/PNG/PDF bytes and valid relationships."""
    word = dict(TABLE_PARTS)
    body = word['word/document.xml']
    start = body.index('<w:tbl>')
    table = body[start:body.index('</w:tbl>') + len('</w:tbl>')]
    word['word/document.xml'] = (
        f'<w:document xmlns:w="{W}"><w:body>'
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/><w:outlineLvl w:val="0"/></w:pPr><w:r><w:t>Quarterly review</w:t></w:r></w:p>'
        '<w:p><w:r><w:t xml:space="preserve">Revenue </w:t></w:r><w:r><w:rPr><w:i/></w:rPr><w:t>up 10%</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Closing paragraph.</w:t></w:r></w:p>' + table +
        '<w:sectPr><w:pgSz w:w="12240" w:h="15840"/></w:sectPr></w:body></w:document>'
    )
    word['word/styles.xml'] = f'<w:styles xmlns:w="{W}"><w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:pPr><w:outlineLvl w:val="0"/></w:pPr></w:style></w:styles>'
    # The styles part this lineage adds joins the settings relationship the fixture already carries,
    # so word/settings.xml stays relationship-owned and Word still opens the copy in the current mode.
    word['word/_rels/document.xml.rels'] = (
        f'<Relationships xmlns="{PKG}"><Relationship Id="styles" Type="{R}/styles" Target="styles.xml"/>'
        f'<Relationship Id="settings" Type="{R}/settings" Target="settings.xml"/></Relationships>'
    )
    word['[Content_Types].xml'] = word['[Content_Types].xml'].replace('</Types>', '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>')
    book = dict(WORKBOOK_PARTS)
    deck = dict(DECK_PARTS)
    # Two existing slides make deletion a normal task, rather than an empty-deck refusal. The
    # copy's relationship id starts past the master, theme and notes-master ids the builder's
    # own PowerPoint-valid chain already occupies.
    deck['ppt/presentation.xml'] = deck['ppt/presentation.xml'].replace('</p:sldIdLst>', '<p:sldId id="257" r:id="rId5"/></p:sldIdLst>')
    deck['ppt/_rels/presentation.xml.rels'] = deck['ppt/_rels/presentation.xml.rels'].replace('</Relationships>', f'<Relationship Id="rId5" Type="{R}/slide" Target="slides/slide2.xml"/></Relationships>')
    deck['ppt/slides/slide2.xml'] = deck['ppt/slides/slide1.xml'].replace('Q3 review', 'Appendix').replace(OLD, 'Unused appendix')
    deck['[Content_Types].xml'] = deck['[Content_Types].xml'].replace('</Types>', '<Override PartName="/ppt/slides/slide2.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/></Types>')
    # The notes slide keeps the notes-master relationship the builder gave it and gains the
    # back reference to its own slide.
    deck['ppt/notesSlides/_rels/notesSlide1.xml.rels'] = deck['ppt/notesSlides/_rels/notesSlide1.xml.rels'].replace('</Relationships>', f'<Relationship Id="rId2" Type="{R}/slide" Target="../slides/slide1.xml"/></Relationships>')
    # A real thumbnail and app metadata exercise byte preservation on ordinary OPC members.
    for parts in (word, book, deck):
        parts['docProps/app.xml'] = '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"><Application>Synthetic comparison</Application></Properties>'
        parts['docProps/thumbnail.png'] = stored_png([[(22,44,66), (88,110,132)]])
        parts['_rels/.rels'] = parts['_rels/.rels'].replace('</Relationships>', f'<Relationship Id="app" Type="{R}/extended-properties" Target="docProps/app.xml"/><Relationship Id="thumb" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail" Target="docProps/thumbnail.png"/></Relationships>')
        parts['[Content_Types].xml'] = parts['[Content_Types].xml'].replace('</Types>', '<Default Extension="png" ContentType="image/png"/><Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/></Types>')
    return dict(docx=stored_zip(word), xlsx=stored_zip(book), pptx=stored_zip(deck), pdf=classic_pdf(PDF_TEXTS))


class OracleFailure(Exception):
    def __init__(self, message, checks=None):
        super().__init__(message)
        self.checks = checks or {}


class Unsupported(Exception):
    pass


class CommandError(Exception):
    pass


def require(condition, message):
    if not condition:
        raise OracleFailure(message)


def xml_shape(node):
    """Namespace-aware semantic tree, ignoring only serializer whitespace and xml:space."""
    if node is None:
        return None
    incidental = {'{http://www.w3.org/XML/1998/namespace}space',
                  '{http://schemas.microsoft.com/office/word/2010/wordml}paraId',
                  '{http://schemas.microsoft.com/office/word/2010/wordml}textId',
                  '{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable'}
    return (node.tag, tuple(sorted((k,v) for k,v in node.attrib.items() if k not in incidental)),
            node.text if node.tag.endswith(('}t', '}delText', '}v', '}f')) else (node.text or '').strip(),
            tuple(xml_shape(c) for c in node))


def text_of(node, namespace=W):
    return ''.join(n.text or '' for n in node.iter(f'{{{namespace}}}t'))


def paragraphs(parts):
    return ET.fromstring(parts['word/document.xml']).find('w:body', NS).findall('w:p', NS)


def run_coverage(paragraph, namespace):
    result = []
    for run in paragraph.findall(f'{{{namespace}}}r'):
        props = xml_shape(run.find(f'{{{namespace}}}rPr'))
        result.extend([(c, props) for c in text_of(run, namespace)])
    return result


def preserved_runs(before, after, namespace, old, new):
    require(text_of(after, namespace) == new, 'replacement text differs')
    b, a = run_coverage(before, namespace), run_coverage(after, namespace)
    require(''.join(c for c,_ in b) == old, 'invalid fixture run coverage')
    require([p for _,p in b] == [p for _,p in a], 'replacement lost or changed run formatting')
    require(xml_shape(before.find(f'{{{namespace}}}pPr')) == xml_shape(after.find(f'{{{namespace}}}pPr')), 'paragraph properties changed')


def rel_target(rels_part, target):
    base = '' if rels_part == '_rels/.rels' else posixpath.dirname(posixpath.dirname(rels_part))
    return posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(base, target))


def relationships_valid(parts):
    for name, data in parts.items():
        if name.endswith(('.xml', '.rels')):
            ET.fromstring(data)
        if not name.endswith('.rels'):
            continue
        ids = set()
        for rel in ET.fromstring(data):
            require(rel.get('Id') not in ids, f'duplicate relationship in {name}')
            ids.add(rel.get('Id'))
            if rel.get('TargetMode') != 'External':
                require(rel_target(name, rel.get('Target', '')) in parts, f'dangling relationship in {name}: {rel.get("Target")}')


def slide_parts(parts):
    rels = 'ppt/_rels/presentation.xml.rels'
    targets = {r.get('Id'): rel_target(rels, r.get('Target')) for r in ET.fromstring(parts[rels])}
    return [targets[n.get(f'{{{R}}}id')] for n in ET.fromstring(parts['ppt/presentation.xml']).findall('p:sldIdLst/p:sldId', NS)]


def cells(parts):
    strings = [text_of(n, S) for n in ET.fromstring(parts.get('xl/sharedStrings.xml', f'<sst xmlns="{S}"/>'.encode())).findall('s:si', NS)]
    result = {}
    for c in ET.fromstring(parts['xl/worksheets/sheet1.xml']).findall('.//s:c', NS):
        kind, value = c.get('t', 'n'), c.findtext('s:v', default='', namespaces=NS)
        formula = c.findtext('s:f', namespaces=NS)
        if formula is not None:
            result[c.get('r')] = ('formula', formula)
        elif kind in ('s','inlineStr','str'):
            result[c.get('r')] = ('string', strings[int(value)] if kind == 's' else text_of(c, S) if kind == 'inlineStr' else value)
        elif kind == 'b':
            result[c.get('r')] = ('boolean', value in ('1','true','TRUE'))
        else:
            result[c.get('r')] = ('number', value)
    return result


def allowed_parts(task, before, after):
    """Task-scoped OPC write sets; existing unlisted members must retain exact bytes."""
    family, op = task['family'], task['operation']
    changed, added, removed = set(), set(), set()
    if family == 'docx':
        changed = {'word/document.xml'}
        if op == 'comment':
            changed |= {'[Content_Types].xml', 'word/_rels/document.xml.rels'}
            added = {'word/comments.xml', 'word/commentsExtended.xml', 'word/commentsIds.xml', 'word/people.xml'}
    elif family == 'xlsx':
        changed = {'xl/worksheets/sheet1.xml', 'xl/workbook.xml'}
        if op in ('string', 'append'):
            changed.add('xl/sharedStrings.xml')
        if op == 'rename':
            changed = {'xl/workbook.xml'}
    elif family == 'pptx':
        changed = {'ppt/notesSlides/notesSlide1.xml'} if op == 'notes' else {'ppt/slides/slide1.xml'}
        if op in ('duplicate','delete'):
            changed = {'[Content_Types].xml', 'ppt/presentation.xml', 'ppt/_rels/presentation.xml.rels'}
            if op == 'delete':
                removed = {'ppt/slides/slide2.xml'}
            else:
                # New part numbers/relationship IDs are not prescribed by the task.
                added = {n for n in after.keys() - before.keys() if re.fullmatch(r'ppt/(slides|notesSlides)/(_rels/)?(slide|notesSlide)\d+\.xml(\.rels)?', n)}
    problems = []
    if after.keys() - before.keys() - added:
        problems.append(f'unexpected added members: {sorted(after.keys() - before.keys() - added)}')
    if before.keys() - after.keys() - removed:
        problems.append(f'unexpected removed members: {sorted(before.keys() - after.keys() - removed)}')
    damaged = [n for n in before.keys() & after.keys() if n not in changed and before[n] != after[n]]
    if damaged: problems.append(f'untouched members changed bytes: {sorted(damaged)}')
    require(not problems, '; '.join(problems))
    return sorted(n for n in before.keys() & after.keys() if n not in changed)


def check_docx(task, before, after):
    op = task['operation']
    b, a = paragraphs(before), paragraphs(after)
    expected = [text_of(p) for p in b]
    if op == 'insert': expected.insert(2, INSERTED)
    elif op == 'delete': del expected[2]
    elif op in ('replace','track'): expected[1] = NEW
    actual = [resolved_text(p, W, accept=True) if op == 'track' else text_of(p) for p in a]
    require(actual == expected, f'body paragraphs differ: {actual}')
    # Remove only the requested target before comparing the rest of the main part.
    br, ar = (ET.fromstring(p['word/document.xml']) for p in (before, after))
    bb, ab = br.find('w:body', NS), ar.find('w:body', NS)
    if op == 'cell':
        bc = bb.findall('w:tbl/w:tr', NS)[1].find('w:tc', NS)
        ac = ab.findall('w:tbl/w:tr', NS)[1].find('w:tc', NS)
        require(text_of(ac) == 'Nova', 'table cell text differs')
        require(len(ac.findall('w:p', NS)) == 1, 'table cell paragraph count changed')
        preserved_runs(bc.find('w:p', NS), ac.find('w:p', NS), W, 'Acme', 'Nova')
        for p in list(bc.findall('w:p', NS)): bc.remove(p)
        for p in list(ac.findall('w:p', NS)): ac.remove(p)
    elif op == 'insert':
        ab.remove(ab.findall('w:p', NS)[2])
    elif op == 'delete':
        bb.remove(bb.findall('w:p', NS)[2])
    else:
        target_before, target_after = bb.findall('w:p', NS)[1], ab.findall('w:p', NS)[1]
        if op == 'replace': preserved_runs(target_before, target_after, W, OLD, NEW)
        if op == 'track':
            require(resolved_text(target_after, W, accept=False) == OLD, 'rejecting native revisions does not restore the original')
            revisions = target_after.findall('.//w:ins', NS) + target_after.findall('.//w:del', NS)
            require(revisions and {n.get(f'{{{W}}}author') for n in revisions} == {AUTHOR}, 'native revision author differs or revisions missing')
        if op == 'comment':
            rels = ET.fromstring(after['word/_rels/document.xml.rels'])
            require(any(r.get('Type') == R+'/comments' and rel_target('word/_rels/document.xml.rels',r.get('Target')) == 'word/comments.xml' for r in rels), 'comments part is not related to the document')
            types = ET.fromstring(after['[Content_Types].xml'])
            require(any(t.get('PartName') == '/word/comments.xml' and t.get('ContentType') == 'application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml' for t in types), 'comments content type missing')
            comments = ET.fromstring(after.get('word/comments.xml', b'<missing/>')).findall('w:comment', NS)
            require(len(comments) == 1 and text_of(comments[0]) == COMMENT, 'comment content differs')
            comment_id = comments[0].get(f'{{{W}}}id')
            require(comments[0].get(f'{{{W}}}author') == AUTHOR, 'comment author differs')
            for tag in ('commentRangeStart','commentRangeEnd','commentReference'):
                marks = ar.findall(f'.//w:{tag}', NS)
                require(len(marks) == 1 and marks[0].get(f'{{{W}}}id') == comment_id and marks[0] in list(target_after.iter()), 'comment not anchored on the target paragraph')
            # Strip comment markers and their empty reference run, then compare the paragraph.
            stripped = copy.deepcopy(target_after)
            for parent in list(stripped.iter()):
                for child in list(parent):
                    if child.tag in {f'{{{W}}}{t}' for t in ('commentRangeStart','commentRangeEnd','commentReference')}:
                        parent.remove(child)
            for run in list(stripped.findall('w:r', NS)):
                if not text_of(run): stripped.remove(run)
            require(xml_shape(stripped) == xml_shape(target_before), 'comment changed the target paragraph')
        bb.remove(target_before); ab.remove(target_after)
    require(xml_shape(br) == xml_shape(ar), 'unrequested DOCX structure or content changed')


def check_xlsx(task, before, after):
    op = task['operation']
    expected = cells(before)
    if op in ('number','string','formula','boolean'):
        expected[task['cell']] = (op, False if op == 'boolean' else task['after'])
    if op == 'append': expected.update(A2=('string','Nova'), B2=('number','12.5'), C2=('boolean',False))
    require(cells(after) == expected, f'cell content/types/formulas differ: {cells(after)}')
    wb, wa = (ET.fromstring(p['xl/workbook.xml']) for p in (before,after))
    if op == 'rename': wb.find('s:sheets/s:sheet', NS).set('name', 'Reviewed')
    # Recalculation policy is allowed, but other workbook content is not.
    for root in (wb,wa):
        for node in root.findall('s:calcPr', NS): root.remove(node)
    require(xml_shape(wb) == xml_shape(wa), 'unrequested workbook metadata changed')
    sb, sa = (ET.fromstring(p['xl/worksheets/sheet1.xml']) for p in (before,after))
    targets = {task.get('cell')} if op in ('number','string','formula','boolean') else {'A2','B2','C2'} if op == 'append' else set()
    for row in sb.findall('s:sheetData/s:row',NS):
        updated = sa.find(f's:sheetData/s:row[@r="{row.get("r")}"]',NS)
        require(updated is not None and row.attrib == updated.attrib, 'existing row properties changed')
    for ref in targets & cells(before).keys():
        bc = sb.find(f'.//s:c[@r="{ref}"]',NS); ac = sa.find(f'.//s:c[@r="{ref}"]',NS)
        require({k:v for k,v in bc.attrib.items() if k != 't'} == {k:v for k,v in ac.attrib.items() if k != 't'}, 'target cell formatting changed')
    for ref in expected.keys() - targets:
        bc = sb.find(f'.//s:c[@r="{ref}"]', NS); ac = sa.find(f'.//s:c[@r="{ref}"]', NS)
        if ref == 'D1' and op == 'number':
            # Literal B1 feeds D1. Refresh 25 or request recalculation, never claim a stale cache current.
            cached = ac.findtext('s:v', namespaces=NS)
            calc = ET.fromstring(after['xl/workbook.xml']).find('s:calcPr', NS)
            require(cached in ('25','25.0') or (calc is not None and calc.get('fullCalcOnLoad') in ('1','true')), 'formula cache stale without a recalculation request')
            for cell in (bc,ac):
                for v in cell.findall('s:v', NS): cell.remove(v)
        require(xml_shape(bc) == xml_shape(ac), f'unrelated cell structure changed: {ref}')
    # Sheet properties outside the value grid may only update the used-range dimension.
    for root in (sb,sa):
        for child in list(root):
            if child.tag in (f'{{{S}}}sheetData', f'{{{S}}}dimension'): root.remove(child)
    require(xml_shape(sb) == xml_shape(sa), 'unrequested sheet metadata changed')
    if before.get('xl/sharedStrings.xml') != after.get('xl/sharedStrings.xml'):
        old = ET.fromstring(before['xl/sharedStrings.xml']).findall('s:si', NS)
        new = ET.fromstring(after['xl/sharedStrings.xml']).findall('s:si', NS)
        require([xml_shape(n) for n in new[:len(old)]] == [xml_shape(n) for n in old], 'existing shared strings changed')


def check_pptx(task, before, after):
    op = task['operation']
    bs, ass = slide_parts(before), slide_parts(after)
    if op in ('duplicate','delete'):
        expected = [before[bs[0]], before[bs[0]], before[bs[1]]] if op == 'duplicate' else [before[bs[0]]]
        def shape(data):
            root = ET.fromstring(data)
            # A duplicate may assign fresh shape IDs within its own slide.
            for index, node in enumerate(root.findall('.//p:cNvPr', NS)): node.set('id',str(index))
            return xml_shape(root)
        require([shape(after[n]) for n in ass] == [shape(d) for d in expected], 'slide order/content/layout differs')
        if op == 'duplicate':
            rels = posixpath.join(posixpath.dirname(ass[1]), '_rels', posixpath.basename(ass[1])+'.rels')
            notes = [r for r in ET.fromstring(after[rels]) if r.get('Type') == R+'/notesSlide']
            require(len(notes) == 1, 'duplicated slide has no notes relationship')
            note_part = rel_target(rels, notes[0].get('Target'))
            require(note_part != 'ppt/notesSlides/notesSlide1.xml', 'duplicated slide shares mutable speaker notes')
            require(xml_shape(ET.fromstring(after[note_part])) == xml_shape(ET.fromstring(before['ppt/notesSlides/notesSlide1.xml'])), 'duplicated notes content differs')
            back_rels = posixpath.join(posixpath.dirname(note_part), '_rels', posixpath.basename(note_part)+'.rels')
            back = [r for r in ET.fromstring(after[back_rels]) if r.get('Type') == R+'/slide']
            require(len(back) == 1 and rel_target(back_rels,back[0].get('Target')) == ass[1], 'duplicated notes point at the wrong slide')
        return
    part = 'ppt/notesSlides/notesSlide1.xml' if op == 'notes' else bs[0]
    b,a = (ET.fromstring(p[part]) for p in (before,after))
    bp, ap = b.findall('.//a:p', NS), a.findall('.//a:p', NS)
    index = 0 if op == 'notes' else 1
    require(len(bp) == len(ap), 'slide/notes paragraph count changed')
    if op == 'replace': preserved_runs(bp[index], ap[index], A, OLD, NEW)
    else: require(text_of(ap[index], A) == task['after'], 'notes text differs')
    for root, target in ((b,bp[index]),(a,ap[index])):
        for parent in root.iter():
            if target in list(parent): parent.remove(target); break
    require(xml_shape(b) == xml_shape(a), 'unrequested slide/notes content or layout changed')


def pdf_page_properties(data):
    """Independently follow the fixture's classic PDF page resources, not a tool read."""
    offsets, trailer = pdf_current_objects(data)
    catalog = pdf_object(data, offsets, pdf_reference(trailer, b'/Root'))
    tree = pdf_object(data, offsets, pdf_reference(catalog, b'/Pages'))
    kids = re.search(rb'/Kids\s*\[([^\]]*)\]', tree)
    require(kids is not None, 'PDF page tree missing')
    result = []
    for ref in re.findall(rb'(\d+)\s+\d+\s+R', kids.group(1)):
        page = pdf_object(data,offsets,int(ref))
        box = re.search(rb'/MediaBox\s*\[([^\]]*)\]',page)
        require(box is not None, 'PDF page geometry missing')
        font = pdf_object(data,offsets,pdf_reference(page,b'/F1'))
        result.append((tuple(box.group(1).split()), re.sub(rb'\s+',b' ',font).strip()))
    return result


def oracle(task, original, produced=None, observed=None):
    try:
        return _oracle(task,original,produced,observed)
    except (ET.ParseError,zipfile.BadZipFile,ValueError,KeyError,IndexError,TypeError,AttributeError) as exc:
        raise OracleFailure(f'malformed or incomplete output: {exc}') from exc


def _oracle(task, original, produced=None, observed=None):
    """Judge only independently decoded input/output bytes and parsed read payloads."""
    op, family = task['operation'], task['family']
    if op in ('preimage','empty_cell'):
        require(produced is None or produced == original, 'unsafe edit changed the file')
        return {'checked': 'unsafe edit left bytes unchanged'}
    if op in ('read','outline'):
        require(produced == original, 'read modified its input')
        parts = package(original) if family != 'pdf' else None
        if family == 'docx':
            if op == 'outline':
                expected = [{'level':1,'text':text_of(paragraphs(parts)[0])}]
            else: expected = text_of(paragraphs(parts)[task['paragraph']])
        elif family == 'xlsx': expected = cells(parts)
        elif family == 'pptx': expected = [text_of(n,A) for n in ET.fromstring(parts[slide_parts(parts)[0]]).findall('.//a:p',NS)]
        else:
            streams = pdf_current_page_texts(original)
            expected = {'count':len(streams), 'texts':[re.search(rb'\((.*?)\) Tj', s).group(1).decode() for s in streams]}
        require(observed == expected, f'read differs: expected {expected!r}, observed {observed!r}')
        return {'checked':'read content/types/outline and unchanged input'}
    require(produced is not None, 'no output file produced')
    if family == 'pdf':
        expected = [pdf_current_page_texts(original)[n-1] for n in task['pages']]
        require(pdf_current_page_texts(produced) == expected, 'PDF page order or untouched content streams differ')
        properties = pdf_page_properties(original)
        require(pdf_page_properties(produced) == [properties[n-1] for n in task['pages']], 'PDF page geometry or font resources changed')
        return {'checked':'independent classic PDF page tree, count, order, exact content streams, geometry and fonts'}
    before, after = package(original), package(produced)
    relationships_valid(after)
    # Provenance and save stamps a tool adds are set aside before the byte rule, and nothing else
    # is: every other untouched member must still keep its exact bytes.
    unstamped_before, unstamped_after, stamps = package_semantics.stamp_filter(before, after)
    # Report both content and byte-preservation discrepancies, so reserialization is visible.
    problems = []
    preserved = []
    checks = {}
    for name, check in (
        ('preservation', lambda: allowed_parts(task,unstamped_before,unstamped_after)),
        ('content', lambda: {'docx':check_docx,'xlsx':check_xlsx,'pptx':check_pptx}[family](task,before,after)),
    ):
        try:
            value = check()
            if isinstance(value,list): preserved = value
            checks[name] = {'pass':True}
        except (OracleFailure, ET.ParseError, KeyError, IndexError) as e:
            problems.append(str(e))
            checks[name] = {'pass':False,'reason':str(e)}
    if stamps: checks['preservation']['stamps'] = stamps
    if problems: raise OracleFailure('; '.join(problems), checks)
    return {'checked':'independent OPC/XML content, formatting and relationships', 'checks':checks, 'byte_identical_members':preserved}


class Runner:
    """All measured tool commands: fixed cwd, 120s deadline, bounded output, process group."""
    def __init__(self, binary, cwd, prefix=(), display=None, timeout=None):
        self.binary, self.cwd, self.prefix = str(binary), Path(cwd), list(prefix)
        self.commands = []
        # None, 'brief' or 'compact'. The flag is injected only into the four verbs that
        # accept it; every other command runs exactly as the default run does.
        self.display = display
        # None keeps the module deadline, so a patched TIMEOUT still governs the default runner.
        self.timeout = timeout

    def run(self, args):
        if self.display and args and args[0] in ('text', 'find', 'inspect', 'edit'):
            index = 2 if args[0] == 'edit' else 1
            args = [*args[:index], '--' + self.display, *args[index:]]
        command = [*self.prefix, self.binary, *map(str,args)]
        index = len(self.commands)
        deadline = TIMEOUT if self.timeout is None else self.timeout
        record = dict(argv=command, timeout_seconds=deadline)
        self.commands.append(record)
        start = time.perf_counter()
        # GenOffice keeps an audit log, its sign-in and its open-tab list under the home directory
        # by default; each points into the task's own folder here, and its documented path policy
        # confines it to that folder, so no run reads or writes anything outside it.
        env = dict(os.environ, OFFICECLI_SKIP_UPDATE='1', OFFICECLI_RESIDENT_FLUSH='each', DOTNET_CLI_TELEMETRY_OPTOUT='1',
                   XDG_CONFIG_HOME=str(self.cwd/'config'), XDG_CACHE_HOME=str(self.cwd/'cache'),
                   TMPDIR=str(self.cwd), DOCX_AUTHOR=AUTHOR, GENOFFICE_AUDIT_LOG='off',
                   GENOFFICE_USER_DATA=str(self.cwd/'genoffice-user-data'), GENOFFICE_AUTH_DIR=str(self.cwd/'genoffice-auth'),
                   GENOFFICE_AI_SETTINGS=str(self.cwd/'genoffice-ai-settings.json'), GENOFFICE_ALLOWED_ROOTS=str(self.cwd))
        failure = None
        with (self.cwd/f'command-{index}.stdout').open('wb') as stdout, (self.cwd/f'command-{index}.stderr').open('wb') as stderr:
            try:
                proc = subprocess.Popen(command, cwd=self.cwd, env=env, stdin=subprocess.DEVNULL,
                                        stdout=stdout, stderr=stderr, start_new_session=True)
            except OSError as exc:
                record.update(returncode=None, error=str(exc), wall_seconds=time.perf_counter()-start, stdout_bytes=0, stderr_bytes=0)
                raise CommandError(str(exc)) from exc
            while proc.poll() is None:
                if time.perf_counter()-start > deadline: failure = 'timeout'
                if os.fstat(stdout.fileno()).st_size > MAX_CAPTURE or os.fstat(stderr.fileno()).st_size > MAX_CAPTURE: failure = 'capture_limit'
                if failure:
                    break
                time.sleep(.005)
            # Kill the group even after the leader exits, preventing background residents.
            try: os.killpg(proc.pid, signal.SIGKILL)
            except ProcessLookupError: pass
            proc.wait()
        stdout_size = (self.cwd/f'command-{index}.stdout').stat().st_size
        stderr_size = (self.cwd/f'command-{index}.stderr').stat().st_size
        with (self.cwd/f'command-{index}.stdout').open('rb') as stream: raw = stream.read(MAX_CAPTURE)
        with (self.cwd/f'command-{index}.stderr').open('rb') as stream: err = stream.read(MAX_CAPTURE)
        record.update(returncode=proc.returncode, wall_seconds=time.perf_counter()-start,
                      stdout_bytes=stdout_size, stderr_bytes=stderr_size, stdout=raw.decode('utf-8',errors='replace'),
                      stderr=err.decode('utf-8',errors='replace'))
        if stdout_size>MAX_CAPTURE or stderr_size>MAX_CAPTURE: failure = 'capture_limit'
        if failure:
            record['error'] = failure
            raise CommandError(failure)
        return record

    def json(self,args):
        rec = self.run(args)
        if rec['returncode'] != 0: raise CommandError(f"command exited {rec['returncode']}: {rec['stderr'][:300]}")
        try: return json.loads(rec['stdout'])
        except ValueError as exc: raise OracleFailure('tool did not return the requested JSON') from exc


def ufo_payload(value):
    # text --format structure prints the structure directly, or in a text receipt.
    if 'result' in value and 'source' in value and 'schema' not in value:
        return value['result']
    return value


def ufo_adapter(task, runner, source, output):
    family,op = task['family'], task['operation']
    args = ['text','--format','structure']
    if family == 'docx':
        args += ['--offset','0','--limit','20'] if op in ('cell','empty_cell') else ['--offset',str(task.get('paragraph',0)),'--limit','1']
    elif family == 'xlsx': args += ['--sheet','Summary','--range','A1:E1']
    elif family == 'pptx': args += ['--slide',str(task.get('slide',1))]
    elif family == 'pdf': args += ['--pages','1,2,3']
    page = ufo_payload(runner.json([*args,str(source)]))
    if op in ('read','outline'):
        if family == 'docx':
            if op == 'outline': return [{'level':h['level'],'text':h['text']} for h in page['outline']]
            return page['paragraphs'][0]['text']
        if family == 'xlsx':
            out = {}
            for c in page['selection']['cells']:
                if c['ref'] == 'E1': continue
                t = c['type']; v = c['formula'].lstrip('=') if t == 'formula' else c['value']
                if t == 'boolean': v = str(v).lower() in ('true','1')
                out[c['ref']] = (t,v)
            return out
        if family == 'pptx': return [p['text'] for s in page['selection']['shapes'] for p in s['paragraphs']]
        return {'count':page['pageCount'], 'texts':[p['text'].strip() for p in page['selection']['items']]}
    digest = page['source']['sha256']
    args = ['edit', 'paragraphs', '--expect-sha256',digest]
    if family == 'docx':
        if op in ('cell','empty_cell'):
            p = next(p for p in page['paragraphs'] if p['table'] == {'table':0,'row':1,'cell':0})
        else: p = page['paragraphs'][0]
        index,preimage = str(p['index']), p['text']
        if op != 'insert': args += ['--expect',f'{index}={task["before"] if op == "preimage" else preimage}']
        if op in ('replace','track','preimage','cell'): args += ['--set',f'{index}={task["after"]}']
        elif op == 'insert': args += ['--insert-after', f'{index}={INSERTED}']
        elif op in ('delete','empty_cell'): args += ['--delete',index]
        elif op == 'comment': args += ['--author',AUTHOR,'--comment',f'{index}={COMMENT}']
        if op == 'track': args += ['--track','--author',AUTHOR]
    elif family == 'xlsx':
        args = ['edit','rows' if op == 'append' else 'sheets' if op == 'rename' else 'cells','--expect-sha256',digest]
        # The September 15 launcher had no sheet rename; `edit sheets` has carried one since.
        if op == 'rename': args += ['--rename',f'{task["sheet"]}={task["after"]}']
        elif op == 'append': args += ['--sheet','Summary','--append-row',json.dumps(task['values'])]
        else:
            c = next(c for c in page['selection']['cells'] if c['ref'] == task['cell'])
            preimage = 'empty' if c['type'] == 'empty' else f'{c["type"]}:{c["value"]}'
            args += ['--expect',f'{task["cell"]}={preimage}','--set',f'{task["cell"]}={op}:{"=" if op == "formula" else ""}{task["after"]}']
    elif family == 'pptx':
        if op in ('duplicate','delete'):
            args = ['edit','slides','--expect-sha256',digest,'--'+op,str(task['slide'])]
            if op == 'duplicate': args += ['--after','1']
        else:
            selection = page['selection']
            if op == 'notes': p = selection['notes']['paragraphs'][0]; idx = 'notes:'+str(p['index'])
            else: p = next(s for s in selection['shapes'] if s['id'] == 3)['paragraphs'][0]; idx = str(p['index'])
            args += ['--slide','1','--expect',f'{idx}={p["text"]}','--set',f'{idx}={task["after"]}']
    else: args = ['edit','pages','--keep','3,1']
    rec = runner.run([*args,'-o',str(output),str(source)])
    if rec['returncode'] not in (0,1): raise CommandError('UFO command usage/runtime error')
    if rec['returncode'] == 1 and op in ('preimage','empty_cell'):
        receipt = json.loads(rec['stdout'])
        if receipt.get('status') != 'refused' or receipt.get('code') != 'parser_refused':
            raise CommandError('UFO failed without a semantic refusal')
    if rec['returncode'] and op not in ('preimage','empty_cell'): raise CommandError('UFO refused a supported task')
    return None


def officecli_adapter(task, runner, source, output, variant=0):
    family,op = task['family'],task['operation']
    if family == 'pdf': raise Unsupported('OfficeCLI documents DOCX/XLSX/PPTX only; no PDF command.')
    # Coordinates default to this suite's pinned fixture; a caller with another file supplies its own.
    path = task.get('path', '/body/p[2]')
    sheet = task.get('sheet', 'Summary')
    read = op in ('read','outline')
    file = str(source if read else output)
    if not read: shutil.copyfile(source,output)
    if read:
        if op == 'outline':
            rec = runner.run(['view',file,'outline','--json']); data = json.loads(rec['stdout'])
            return office_outline(data)
        if family == 'docx': return office_data(runner.json(['get',file,path,'--json']))['text']
        if family == 'xlsx':
            data = office_data(runner.json(['get',file,'/Summary/A1:D1','--json']))
            result = {}
            for c in walk_nodes(data):
                if c.get('type','').lower() != 'cell': continue
                props = c.get('format',{}); kind = props.get('type','Number').lower()
                if kind in ('sharedstring','inlinestring','str'): kind = 'string'
                value = c.get('text')
                if props.get('formula') is not None: kind,value = 'formula',props['formula'].lstrip('=')
                if kind == 'boolean': value = str(value).lower() in ('true','1')
                result[c['path'].rsplit('/',1)[1]] = (kind,value)
            return result
        data = office_data(runner.json(['get',file,'/slide[1]','--depth','1','--json']))
        return [c['text'] for c in data.get('children',[]) if 'text' in c]
    if family == 'docx':
        if op in ('replace','track','preimage'):
            args = ['set',file,path,'--find',task['before'],'--replace',task['after']]
            if op == 'replace' and variant == 1:
                args = ['set',file,path,'--find','10','--replace','12']
            if op == 'track': args += ['--prop','revision.author='+task.get('author',AUTHOR)]
        elif op == 'insert': args = ['add',file,'/body','--type','paragraph','--after',path,'--prop','text='+task.get('text',INSERTED)]
        elif op == 'delete': args = ['remove',file,task.get('deletePath','/body/p[3]')]
        elif op == 'comment': args = ['add',file,path,'--type','comment','--prop','text='+task.get('text',COMMENT),'--prop','author='+task.get('author',AUTHOR)]
        elif op == 'cell': args = ['set',file,'/body/tbl[1]/tr[2]/tc[1]/p[1]','--find','Acme','--replace','Nova']
        else: args = ['remove',file,'/body/tbl[1]/tr[2]/tc[1]/p[1]']
    elif family == 'xlsx':
        if op == 'rename': args = ['set',file,'/'+sheet,'--prop','name='+task['after']]
        elif op == 'append':
            row = str(task.get('row',2))
            commands = [dict(command='set',path=f'/{sheet}/{c}{row}',props={'value':v.split(':',1)[1],'type':v.split(':',1)[0]}) for c,v in task['values'].items()]
            args = ['batch',file,'--commands',json.dumps(commands)]
        else:
            args = ['set',file,f'/{sheet}/'+task['cell'],'--prop',('formula=' if op == 'formula' else 'value=')+task['after']]
            if op != 'formula': args += ['--prop','type='+op]
    else:
        if op == 'replace':
            # The public shape schema is ambiguous about multi-run text preservation.
            # Two predeclared attempts: whole shape text, then the documented run setter.
            slide, shape = task.get('slide',1), task.get('shapeOrdinal',2)
            target = f'/slide[{slide}]/shape[{shape}]' if variant == 0 else f'/slide[{slide}]/shape[{shape}]/paragraph[1]/run[2]'
            args = ['set',file,target,'--prop','text='+(task['after'] if variant == 0 else task.get('changedSpan','up 12%'))]
        elif op == 'notes': args = ['set',file,f"/slide[{task.get('slide',1)}]/notes",'--prop','text='+task['after']]
        elif op == 'duplicate': args = ['add',file,'/','--type','slide','--from',f"/slide[{task.get('slide',1)}]",'--index',str(task.get('slide',1))]
        else: args = ['remove',file,f"/slide[{task.get('slide',2)}]"]
    rec = runner.run([*args,'--json'])
    if rec['returncode']: raise CommandError(f'OfficeCLI exited {rec["returncode"]}')
    return None


def office_data(value):
    value = value.get('data',value)
    if isinstance(value,dict) and value.get('matches') == 1:
        return value['results'][0]
    return value


def walk_nodes(node):
    if isinstance(node,list):
        for child in node: yield from walk_nodes(child)
    elif isinstance(node,dict):
        yield node
        for child in node.get('children',[]): yield from walk_nodes(child)


def office_outline(data):
    data = office_data(data)
    if isinstance(data,str):
        return [{'level':int(m.group(1)),'text':m.group(2)} for m in re.finditer(r'\[H(\d)\]\s*(.+)', data)]
    return [{'level':n['level'],'text':n['text']} for n in data.get('headings',[]) ]


def docx_adapter(task, runner, source, output, variant=0):
    if task['family'] != 'docx': raise Unsupported('docx-cli is DOCX only.')
    op = task['operation']; file = str(source)
    if op == 'read':
        rec = runner.run(['read',file,'--from',task.get('at','p1'),'--to',task.get('at','p1')])
        if rec['returncode']: raise CommandError('docx-cli read failed')
        # This fixture has ordinary plain/italic runs, no literal Markdown punctuation.
        value = re.sub(r'<!--.*?-->', '', rec['stdout'], flags=re.S).strip()
        return value.replace('*','').replace('_','')
    if op == 'outline':
        value = runner.json(['outline',file,'--json'])
        return docx_outline(value)
    # The paragraph address defaults to this suite's pinned fixture; another file supplies its own.
    at = task.get('at','p1')
    if op in ('replace','track','preimage'):
        args = ['replace',file,task['before'],task['after'],'--at',at,'--exact']
        if op == 'replace' and variant == 1:
            args = ['replace',file,'10','12','--at',at,'--exact']
        if op == 'track': args += ['--track']
    elif op == 'insert': args = ['insert',file,'--after',at,'--text',task.get('text',INSERTED)]
    elif op == 'delete': args = ['delete',file,'--at',task.get('deleteAt','p2')]
    elif op == 'comment': args = ['comments','add',file,'--at',at,'--text',task.get('text',COMMENT),'--author',task.get('author',AUTHOR)]
    elif op == 'cell': args = ['replace',file,'Acme','Nova','--at','t0:r1c0:p0','--exact']
    else: args = ['delete',file,'--at','t0:r1c0:p0']
    rec = runner.run([*args,'-o',str(output)])
    if rec['returncode']:
        refusal = json.loads(rec['stdout']) if op == 'preimage' else {}
        if refusal.get('code') != 'MATCH_NOT_FOUND': raise CommandError(f'docx-cli exited {rec["returncode"]}')
    return None


def docx_outline(value):
    entries = value if isinstance(value,list) else value.get('headings',[])
    return [{'level':n['level'],'text':n['text']} for n in entries]


def genoffice_result(runner, args):
    """One GenOffice command with --json; its single result object, or a CommandError naming its reason."""
    record = runner.run([*args, '--json'])
    try:
        value = json.loads(record['stdout'])
    except ValueError:
        value = None
    if record['returncode'] or not isinstance(value, dict) or value.get('status') not in ('ok', 'partial'):
        reason = value.get('error') if isinstance(value, dict) else None
        message = (value.get('message') if isinstance(value, dict) else None) or record['stderr'][:200]
        raise CommandError(f'GenOffice exited {record["returncode"]}: {reason or "no JSON result"}: {str(message)[:200]}')
    if value.get('status') == 'partial':
        raise CommandError('GenOffice applied part of the batch')
    return value.get('detail') or {}


def genoffice_apply(runner, domain, source, output, ops, *options):
    """`genoffice <domain> apply` with the ops in a file beside the task, publishing to [output]."""
    path = runner.cwd / f'ops-{len(runner.commands)}.json'
    path.write_text(json.dumps(ops), encoding='utf-8')
    return genoffice_result(runner, [domain, 'apply', str(source), '--ops', str(path), '--out', str(output), *options])


def genoffice_adapter(task, runner, source, output, variant=0):
    """GenOffice's documented ops: `docs`, `sheet` and `slides` read and apply, and `info`/`convert` for PDF.

    Every write names the source and publishes with --out, so the source copy stays as it was. Where
    a whole-text find does not match across differently formatted runs, the declared second attempt
    finds only the changed characters; for the table cell it reads the table as HTML and rewrites
    that one block, the route GenOffice's own guide gives for table text. Every attempt counts.
    """
    family, op = task['family'], task['operation']
    file = str(source)
    if family == 'pdf':
        if op != 'read':
            raise Unsupported('GenOffice 0.10.1038 has no command that edits PDF pages; its PDF editor is in the app window.')
        # The command line reads a PDF's text by converting it to Word locally and reading that.
        count = genoffice_result(runner, ['info', file])['pages']
        converted = runner.cwd / 'converted.docx'
        genoffice_result(runner, ['convert', file, '--to', 'docx', '--out', str(converted)])
        items = genoffice_result(runner, ['docs', 'read', str(converted), '--full'])['items']
        return {'count': count, 'texts': [i['text'].strip() for i in items if (i.get('text') or '').strip()]}
    if family == 'docx':
        if op == 'empty_cell':
            raise Unsupported('GenOffice docs ops address whole body blocks; a table is one block and no op '
                              'removes one paragraph from a table cell.')
        blocks = genoffice_result(runner, ['docs', 'read', file, '--full'])['items']
        # The task counts the body's own paragraphs; GenOffice counts every body block, tables included.
        paragraph_blocks = [b['index'] for b in blocks if b.get('type') != 'table']
        if op == 'outline':
            return [{'level': b['level'], 'text': b['text']} for b in blocks if b.get('type') == 'heading']
        if op == 'read':
            index = paragraph_blocks[task['paragraph']]
            return next(b['text'] for b in blocks if b['index'] == index)
        block = paragraph_blocks[task.get('paragraph', 0)]
        options = ['--track', '--author', task.get('author', AUTHOR)] if op == 'track' else []
        if op in ('replace', 'track', 'preimage'):
            find, replace = (task['before'], task['after']) if variant == 0 or op == 'preimage' else ('10', '12')
            ops = [{'op': 'findReplace', 'find': find, 'replace': replace, 'matchCase': True, 'target': {'blockIndexes': [block]}}]
        elif op == 'insert':
            ops = [{'op': 'insert_content', 'html': f'<p>{task.get("text", INSERTED)}</p>', 'afterBlockIndex': block}]
        elif op == 'delete':
            ops = [{'op': 'deleteBlocks', 'target': {'blockIndexes': [paragraph_blocks[task['paragraph']]]}}]
        elif op == 'comment':
            ops = [{'op': 'add_comment', 'blockIndex': block, 'comment': task.get('text', COMMENT), 'author': task.get('author', AUTHOR)}]
        else:
            table = [b['index'] for b in blocks if b.get('type') == 'table'][task['table']]
            if variant == 0:
                ops = [{'op': 'findReplace', 'find': task['before'], 'replace': task['after'], 'matchCase': True,
                        'target': {'blockIndexes': [table]}}]
            else:
                html = genoffice_result(runner, ['docs', 'read', file, '--range', f'{table}-{table}', '--html'])['html']
                if html.count(f'<td>{task["before"]}</td>') != 1:
                    raise CommandError('the table HTML GenOffice read back does not hold the cell text exactly once')
                ops = [{'op': 'replace_blocks', 'startBlockIndex': table, 'endBlockIndex': table,
                        'html': html.replace(f'<td>{task["before"]}</td>', f'<td>{task["after"]}</td>')}]
        try:
            genoffice_apply(runner, 'docs', source, output, ops, *options)
        except CommandError as exc:
            # A find that no longer matches is GenOffice's refusal of a stale preimage: nothing is written.
            if op == 'preimage' and 'target_not_found' in str(exc) and not output.exists():
                return None
            raise
        return None
    if family == 'xlsx':
        sheet = task.get('sheet', 'Summary')
        if op == 'read':
            detail = genoffice_result(runner, ['sheet', 'read', file, '--sheet', sheet, '--range', task['range']])
            first = re.match(r'([A-Z]+)(\d+)', task['range'].split(':')[0])
            result = {}
            for r, row in enumerate(detail['rows']):
                for c, value in enumerate(row):
                    if value is None or value == '':
                        continue
                    ref = f'{column_name(column_number(first.group(1)) + c)}{int(first.group(2)) + r}'
                    formula = detail.get('formulas', {}).get(ref)
                    if formula is not None: result[ref] = ('formula', formula.lstrip('='))
                    elif isinstance(value, bool): result[ref] = ('boolean', value)
                    elif isinstance(value, (int, float)): result[ref] = ('number', str(int(value)) if float(value).is_integer() else repr(value))
                    else: result[ref] = ('string', str(value))
            return result
        if op == 'rename':
            ops = [{'op': 'rename_sheet', 'sheet': sheet, 'name': task['after']}]
        elif op == 'append':
            last = genoffice_result(runner, ['sheet', 'read', file, '--sheet', sheet])['range'].split(':')[-1]
            row = int(re.search(r'\d+', last).group(0)) + 1
            kinds = {'string': str, 'number': float, 'boolean': lambda v: v.lower() == 'true'}
            values = [kinds[v.split(':', 1)[0]](v.split(':', 1)[1]) for _, v in sorted(task['values'].items())]
            ops = [{'op': 'set_range', 'sheet': sheet, 'start': f'A{row}', 'values': [values]}]
        else:
            # The cell's current value is the preimage GenOffice checks before it writes.
            before = {'number': lambda v: float(v), 'boolean': lambda v: v.upper() == 'TRUE'}.get(op, lambda v: v)
            expected = None if task['before'] is None else before(task['before'])
            if op == 'formula':
                ops = [{'op': 'set_formula', 'sheet': sheet, 'address': task['cell'], 'formula': '=' + task['after']}]
            else:
                value = {'number': float, 'boolean': lambda v: v.lower() == 'true'}.get(op, str)(task['after'])
                ops = [{'op': 'set_cell', 'sheet': sheet, 'address': task['cell'], 'value': value, 'expectedValue': expected}]
        genoffice_apply(runner, 'sheet', source, output, ops)
        return None
    slide = task.get('slide', 1) - 1
    if op == 'read':
        page = genoffice_result(runner, ['slides', 'read', file, '--full'])['pages'][slide]
        return [line for e in page['elements'] if 'text' in e for line in e['text'].split('\n')]
    if op == 'replace':
        find, replace = (task['before'], task['after']) if variant == 0 else ('10', '12')
        ops = [{'op': 'findReplace', 'find': find, 'replace': replace, 'matchCase': True, 'slideIndex': slide}]
    elif op == 'notes':
        ops = [{'op': 'setNotes', 'target': {'slide': slide}, 'text': task['after']}]
    elif op == 'duplicate':
        # The copy lands immediately after the source, which is where the task asks for it.
        ops = [{'op': 'duplicateSlide', 'target': {'slide': slide}}]
    else:
        ops = [{'op': 'deleteSlide', 'target': {'slide': slide}}]
    genoffice_apply(runner, 'slides', source, output, ops)
    return None


def network_mode(cwd):
    """Probe capability, not merely executable presence; never require privileges."""
    executable = shutil.which('unshare')
    attempts = []
    if executable:
        for flags in (['-n'], ['-Urn']):
            r = Runner(executable,cwd).run([*flags,'--','true'])
            attempts.append(r)
            if r['returncode'] == 0:
                return [executable,*flags,'--'], {'disabled':True,'method':'unshare '+ ' '.join(flags),'probes':attempts}
    return [], {'disabled':False,'method':'unshare unavailable or denied; downloads precede measurement; tasks have no network need; OfficeCLI updates disabled','probes':attempts}


def pinned_download(name, pin):
    """The release asset's bytes, only once they hash to the pin. Never executes an installer."""
    with urllib.request.urlopen(pin['url'],timeout=120) as response:
        data = response.read(MAX_DOWNLOAD+1)
    if len(data)>MAX_DOWNLOAD or sha(data) != pin['sha256']: raise ValueError(f'{name}: release binary SHA-256 mismatch')
    return data


def deb_data_archive(data):
    """The `data.tar.*` member of a Debian package, which is an ar archive of three members."""
    if not data.startswith(b'!<arch>\n'): raise ValueError('not a Debian package: no ar signature')
    at = 8
    while at + 60 <= len(data):
        header = data[at:at+60]
        member = header[:16].decode('ascii','replace').strip().rstrip('/')
        size = int(header[48:58].decode('ascii').strip())
        if member.startswith('data.tar'):
            return data[at+60:at+60+size]
        at += 60 + size + (size & 1)
    raise ValueError('the Debian package carries no data archive')


def provision_deb(cache, name, pin, download):
    """A pinned Debian package unpacked in place under the cache, never installed.

    The package is kept and hash-checked like a binary. Its files are unpacked with tar's `data`
    filter (no absolute paths, no links out of the tree, no device or set-id bits), and the files
    the run executes are hashed at unpacking and checked on every later run, so a changed launcher
    is unpacked again from the verified package rather than run.
    """
    archive = cache/pin['asset']
    root = cache/f"{name}-{pin['tag'].lstrip('v')}"
    marker = root/'.pinned.json'
    if not (archive.is_file() and sha(archive.read_bytes()) == pin['sha256']):
        if not download: raise ValueError(f'{name} is missing or differs from its pinned SHA-256; run compare --download')
        archive.write_bytes(pinned_download(name, pin))
    try:
        recorded = json.loads(marker.read_text(encoding='utf-8'))
    except (OSError, ValueError):
        recorded = None
    current = recorded is not None and recorded.get('sha256') == pin['sha256'] and all(
        (root/member).is_file() and sha((root/member).read_bytes()) == recorded.get('members', {}).get(member)
        for member in pin['verify'])
    if not current:
        staging = Path(tempfile.mkdtemp(prefix=f'.{name}-', dir=cache))
        try:
            with tarfile.open(fileobj=io.BytesIO(deb_data_archive(archive.read_bytes())), mode='r:*') as tar:
                tar.extractall(staging, filter='data')
            members = {member: sha((staging/member).read_bytes()) for member in pin['verify']}
            (staging/'.pinned.json').write_bytes(json_bytes(dict(sha256=pin['sha256'], members=members)))
            if root.exists(): shutil.rmtree(root)
            staging.rename(root)
        finally:
            if staging.exists(): shutil.rmtree(staging)
        recorded = json.loads(marker.read_text(encoding='utf-8'))
    return dict(pin, path=str(root/pin['entry']), unpacked=str(root), members=recorded['members'])


def provision(cache, download=False):
    cache.mkdir(parents=True,exist_ok=True)
    result = {}
    for name,pin in PINS.items():
        if pin.get('package') == 'deb':
            result[name] = provision_deb(cache, name, pin, download); continue
        binary = cache/name
        if binary.is_file() and sha(binary.read_bytes()) == pin['sha256']:
            result[name] = dict(pin,path=str(binary)); continue
        if not download: raise ValueError(f'{name} is missing or differs from its pinned SHA-256; run compare --download')
        # Save downloaded bytes only after checking the hard pin. Never execute installers.
        data = pinned_download(name, pin)
        binary.write_bytes(data); binary.chmod(0o755)
        (cache/(name+'.json')).write_bytes(json_bytes(pin))
        result[name] = dict(pin,path=str(binary))
    return result


def percentile95(values):
    return sorted(values)[math.ceil(.95*len(values))-1] if values else None


def summarize(records, identities):
    summary = {}
    for tool,identity in identities.items():
        rows = [r for r in records if r['tool'] == tool]
        counts = Counter(r['result'] for r in rows)
        # Unsupported zero-command tasks are excluded from latency to avoid rewarding narrow tools.
        times = [r['wall_seconds'] for r in rows if r['commands']]
        summary[tool] = dict(identity, **{k:counts[k] for k in ('pass','wrong','unsupported','error')},
                             measured_tasks=len(times), median_seconds=statistics.median(times) if times else None,
                             p95_seconds=percentile95(times), stdout_bytes=sum(r['stdout_bytes'] for r in rows),
                             command_count=sum(r['command_count'] for r in rows))
    return summary


# Tasks where a tool's documentation left two readings, so the adapter declares a second attempt in
# advance. Both attempts run and are counted. GenOffice's find does not match text that spans two
# differently formatted runs, and it does not reach table cells, so each of those gets its guide's
# other route; see genoffice_adapter.
SYNTHETIC_ALTERNATIVES = {
    'officecli': ('pptx_replace', 'docx_replace'),
    'docx-cli': ('docx_replace',),
    'genoffice': ('docx_replace', 'docx_track', 'docx_cell', 'pptx_replace'),
}


def execute_task(task, name, binary, data, prefix, receipt_root, ufo_display=None):
    start = time.perf_counter()
    record = dict(schema_version=1,task=task,tool=name,result='error',attempts=[],commands=[])
    with tempfile.TemporaryDirectory(prefix='task-',dir=receipt_root) as temporary:
        cwd = Path(temporary); source = cwd/('input.'+task['family']); source.write_bytes(data)
        output = cwd/('output.'+task['family']); runner = Runner(binary,cwd,prefix,display=ufo_display if name == 'ufo' else None)
        try:
            alternatives = task['id'] in SYNTHETIC_ALTERNATIVES.get(name, ())
            for variant in range(2 if alternatives else 1):
                if output.exists(): output.unlink()
                try:
                    if name == 'ufo': observed = ufo_adapter(task,runner,source,output)
                    elif name == 'officecli': observed = officecli_adapter(task,runner,source,output,variant)
                    elif name == 'genoffice': observed = genoffice_adapter(task,runner,source,output,variant)
                    else: observed = docx_adapter(task,runner,source,output,variant)
                    require(source.read_bytes() == data, 'tool changed the immutable source')
                    produced = source.read_bytes() if task['operation'] in ('read','outline') else output.read_bytes() if output.exists() else None
                    check = oracle(task,data,produced,observed)
                    record['attempts'].append(dict(variant=variant,result='pass',oracle=check,observed=observed))
                    record.update(result='pass',selected_variant=variant,oracle=check)
                    record.pop('reason',None)
                    record.pop('checks',None)
                except (OracleFailure, ET.ParseError, zipfile.BadZipFile) as exc:
                    checks = getattr(exc,'checks',{})
                    record['attempts'].append(dict(variant=variant,result='wrong',reason=str(exc),checks=checks))
                    record.update(result='wrong',reason=str(exc),checks=checks)
                except CommandError as exc:
                    record['attempts'].append(dict(variant=variant,result='error',reason=str(exc)))
                    record.update(result='error',reason=str(exc))
                    if source.read_bytes() != data:
                        record.update(result='wrong', reason='failed command modified its source')
                    elif output.exists() and output.read_bytes() != data:
                        try: oracle(task,data,output.read_bytes())
                        except (OracleFailure,ET.ParseError,zipfile.BadZipFile) as error:
                            record.update(result='wrong',reason=f'failed command produced an invalid file: {error}')
                # Keep both attempted output files when documentation required a fallback.
                if output.exists():
                    saved = receipt_root/(task['id']+'-'+name+f'-attempt-{variant}.'+task['family'])
                    shutil.copyfile(output,saved)
                    record['attempts'][-1]['output'] = dict(path=saved.name,sha256=sha(output.read_bytes()))
                if record['result'] == 'pass':
                    break
            if output.exists():
                artifact = receipt_root/(task['id']+'-'+name+'.'+task['family'])
                shutil.copyfile(output,artifact)
                record['output'] = dict(path=artifact.name,sha256=sha(output.read_bytes()),bytes=output.stat().st_size)
        except Unsupported as exc: record.update(result='unsupported',reason=str(exc))
        except (KeyError,IndexError,TypeError,ValueError,StopIteration,OSError) as exc:
            record.update(result='error',reason=f'adapter/runtime: {type(exc).__name__}: {exc}')
        record.update(commands=runner.commands, command_count=len(runner.commands),
                      stdout_bytes=sum(c['stdout_bytes'] for c in runner.commands),
                      wall_seconds=time.perf_counter()-start, source_sha256=sha(data))
    (receipt_root/(task['id']+'-'+name+'.json')).write_bytes(json_bytes(record))
    return record


def markdown_table(summary):
    rows = ['| Tool | Pass | Wrong | Unsupported | Error | Median s | p95 s | Commands | Stdout bytes |',
            '| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: |']
    for tool,r in summary.items():
        med = f'{r["median_seconds"]:.3f}' if r['median_seconds'] is not None else 'n/a'
        p95 = f'{r["p95_seconds"]:.3f}' if r['p95_seconds'] is not None else 'n/a'
        rows.append(f'| {tool} | {r["pass"]} | {r["wrong"]} | {r["unsupported"]} | {r["error"]} | {med} | {p95} | {r["command_count"]} | {r["stdout_bytes"]} |')
    return '\n'.join(rows)


# ---------------------------------------------------------------- real-world corpus mode

CORPUS_VERSION = 1
CORPUS_MAX_BYTES = 16 * 1024 * 1024
CORPUS_MAX_PACKAGE = 128 * 1024 * 1024
CORPUS_TIMEOUT = 90
CORPUS_FAMILIES = ('docx', 'xlsx', 'pptx')
CORPUS_TOOLS = ('ufo', 'officecli', 'genoffice', 'docx-cli', 'python-docx', 'openpyxl', 'python-pptx')
CORPUS_RESULTS = ('pass', 'wrong', 'declined', 'unsupported', 'error')
WORD_PATTERN = re.compile(r"[A-Za-z][A-Za-z0-9'-]{2,20}")
REPLACEMENTS = ('Reviewed', 'Amended', 'Revised', 'Updated', 'Checked')
APPENDED_PARAGRAPH = 'Appended by the corpus comparison.'
APPENDED_ROW = 'Corpus comparison row'
WRITTEN_NUMBER = '12.5'
# Pinned before measurement, hashes taken from the downloaded wheels. --no-deps plus an
# explicit dependency list keeps the resolver out of the measured environment.
LIBRARY_PINS = (
    ('python-docx', '1.1.2', '08c20d6058916fb19853fcf080f7f42b6270d89eac9fa5f8c15f691c0017fabe'),
    ('openpyxl', '3.1.5', '5282c12b107bffeef825f4617dc029afaf41d0ea60823bbb665ef3079dc79de2'),
    ('python-pptx', '1.0.2', '160838e0b8565a8b1f67947675886e9fea18aa5e795db7ae531606d68e785cba'),
    ('lxml', '6.1.3', '527195c188d7d0af748cd48d220ab8cdc5cb99be3d49ac4d9be7324d8abf9bc0'),
    ('Pillow', '12.3.0', '23d27a3e0307ec2244cc51e7287b919aa68d097504ebe19df4e76a98a3eea5bd'),
    ('XlsxWriter', '3.2.9', '9a5db42bc5dff014806c58a20b9eae7322a134abb6fce3c92c181bfb275ec5b3'),
    ('et_xmlfile', '2.0.0', '7a91720bc756843502c3b7504c77b8fe44217c85c537d85037f0f536151b2caa'),
    ('typing_extensions', '4.16.0', '481caa481374e813c1b176ada14e97f1f67a4539ce9cfeb3f350d78d6370c2e8'),
)
LIBRARY_FAMILY = {'python-docx': 'docx', 'openpyxl': 'xlsx', 'python-pptx': 'pptx'}
LIBRARY_LICENSE = {'python-docx': 'MIT', 'openpyxl': 'MIT', 'python-pptx': 'MIT'}


def corpus_package(data):
    return package(data, CORPUS_MAX_PACKAGE)


def element_text(node, namespace):
    """Every text character the element carries, in document order."""
    return ''.join(n.text or '' for n in node.iter(f'{{{namespace}}}t'))


def docx_body(parts):
    body = ET.fromstring(parts['word/document.xml']).find(f'{{{W}}}body')
    require(body is not None, 'the main document part has no body')
    return body


def docx_nodes(body):
    """Body paragraphs in document order with their table flag, as a paragraph editor addresses them."""
    found = []

    def walk(node, in_table, depth):
        if depth > 32:
            return
        for child in node:
            if child.tag == f'{{{W}}}p':
                found.append((child, in_table))
            elif child.tag == f'{{{W}}}tbl':
                for row in child.findall(f'{{{W}}}tr'):
                    for cell in row.findall(f'{{{W}}}tc'):
                        walk(cell, True, depth + 1)
            elif child.tag == f'{{{W}}}sdt':
                content = child.find(f'{{{W}}}sdtContent')
                if content is not None:
                    walk(content, in_table, depth + 1)
    walk(body, False, 0)
    return found


def plain_docx_paragraph(node):
    """Only paragraph properties and plain text runs: what every tool under test can rewrite."""
    for child in node:
        if child.tag == f'{{{W}}}pPr':
            continue
        if child.tag != f'{{{W}}}r':
            return False
        for grandchild in child:
            if grandchild.tag not in (f'{{{W}}}rPr', f'{{{W}}}t'):
                return False
    return True


def plain_slide_paragraph(node):
    for child in node:
        if child.tag not in (f'{{{A}}}pPr', f'{{{A}}}r', f'{{{A}}}endParaRPr'):
            return False
    for run in node.findall(f'{{{A}}}r'):
        for grandchild in run:
            if grandchild.tag not in (f'{{{A}}}rPr', f'{{{A}}}t'):
                return False
    return True


def other_part_bytes(parts, keep):
    """Lowercased bytes of every member except the one the task edits, for word uniqueness."""
    return b'\n'.join(data.lower() for name, data in sorted(parts.items()) if name != keep)


def unique_word(run_text, holder_texts, elsewhere):
    """A word of the file's own text that occurs once in its part and nowhere else in the package.

    Uniqueness is what makes one task statement mean the same edit to every tool: a scoped
    find, a whole-file find and a coordinate all land on the same characters.

    It has to hold of the part's READABLE TEXT, not only of its text nodes one at a time. A
    producer may split a word across runs, so a word that sits inside exactly one `<a:t>` can
    still occur several times in the text those nodes spell out, and then the request names more
    than one edit: the content oracle reads the part as running text and would insist on the
    first occurrence, while the residual oracle insists on the one node that carries the word.
    A word that is not unique under both readings states no task and is passed over.
    """
    running = ''.join(holder_texts).lower()
    for word in WORD_PATTERN.findall(run_text):
        lowered = word.lower()
        if sum(text.lower().count(lowered) for text in holder_texts) != 1:
            continue
        if running.count(lowered) != 1:
            continue
        if lowered.encode() in elsewhere:
            continue
        return word
    return None


def free_replacement(taken, word=''):
    """A neutral replacement the file does not carry and that does not overlap the word itself.

    Overlap matters: a replacement holding the word would let a repeated find/replace match its
    own output, so the task statement would no longer name one edit.
    """
    lowered = word.lower()
    for suffix in range(100):
        for base in REPLACEMENTS:
            candidate = base + (str(suffix) if suffix else '')
            low = candidate.lower()
            if lowered and (lowered in low or low in lowered):
                continue
            if low.encode() not in taken:
                return candidate
    return None


def column_number(letters):
    value = 0
    for letter in letters:
        value = value * 26 + (ord(letter) - 64)
    return value


def column_name(number):
    letters = ''
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def sheet_rows(root):
    """Every sheetData row with its resolved number.

    A row may omit its own `r`, and a file in the wild may carry one that is not a worksheet
    row number at all (`r="0"`); either way the row is still at the next position down.
    """
    grid = root.find('s:sheetData', NS)
    number = 0
    for row in (list(grid) if grid is not None else []):
        if row.tag != f'{{{S}}}row':
            continue
        declared = row.get('r') or ''
        number = int(declared) if declared.isdigit() and int(declared) >= 1 else number + 1
        yield number, row


def sheet_grid(data):
    """Every cell with its resolved address: a row or a cell may omit its own `r` attribute.

    SpreadsheetML lets a writer leave the address implied by position, and real files do
    (`<c s="2" t="s">` after `<c r="A1">` is B1). Reading those cells as unaddressed would
    both mis-measure a sheet's used range and lose the cell.
    """
    for number, row in sheet_rows(ET.fromstring(data)):
        column = 0
        for cell in row.findall('s:c', NS):
            reference = cell.get('r') or ''
            if re.fullmatch(r'[A-Z]+[0-9]+', reference):
                column = column_number(re.match(r'[A-Z]+', reference).group(0))
            else:
                column += 1
                reference = f'{column_name(column)}{number}'
            yield reference, cell, number, column


def workbook_sheets(parts):
    """Sheet name, part and visibility from the workbook and its relationships, not from a tool."""
    rels = 'xl/_rels/workbook.xml.rels'
    targets = {r.get('Id'): rel_target(rels, r.get('Target', '')) for r in ET.fromstring(parts[rels])}
    found = []
    for sheet in ET.fromstring(parts['xl/workbook.xml']).findall('s:sheets/s:sheet', NS):
        part = targets.get(sheet.get(f'{{{R}}}id'))
        found.append((sheet.get('name'), part, sheet.get('state') or 'visible'))
    return found


def plan_docx(parts):
    """Two tasks a Word file can state about itself: one word rewrite, one appended paragraph."""
    tasks, skipped = [], []
    body = docx_body(parts)
    nodes = docx_nodes(body)
    direct = [node for node in body if node.tag == f'{{{W}}}p']
    holders = [n.text or '' for n in ET.fromstring(parts['word/document.xml']).iter()
               if n.tag in (f'{{{W}}}t', f'{{{W}}}instrText', f'{{{W}}}delText')]
    elsewhere = other_part_bytes(parts, 'word/document.xml')
    taken = elsewhere + b'\n' + '\n'.join(holders).lower().encode()
    chosen, ordinal = None, 0
    for index, (node, in_table) in enumerate(nodes):
        if not in_table and plain_docx_paragraph(node):
            for run in node.findall(f'{{{W}}}r'):
                for text_node in run.findall(f'{{{W}}}t'):
                    word = unique_word(text_node.text or '', holders, elsewhere)
                    replacement = free_replacement(taken, word or '') if word else None
                    if replacement is not None:
                        chosen = (index, ordinal, node, word, replacement)
                        break
                if chosen:
                    break
        if chosen:
            break
        if not in_table:
            ordinal += 1
    if chosen is None:
        skipped.append('no plain body paragraph holds a word that occurs exactly once in the package')
    else:
        index, ordinal, node, word, replacement = chosen
        tasks.append(dict(
            operation='docx_replace_word', family='docx',
            request=f'Replace the word "{word}" with "{replacement}". Change nothing else.',
            word=word, replacement=replacement, paragraphIndex=index, paragraphOrdinal=ordinal,
            paragraphPath=f'/body/p[{ordinal + 1}]', locator=f'p{ordinal}',
            paragraphText=element_text(node, W)))
    if not direct:
        skipped.append('the body has no paragraph of its own to append after')
    elif nodes and nodes[-1][0] is not direct[-1]:
        skipped.append('the last paragraph of the body sits inside a table, so "at the end" is ambiguous')
    else:
        tasks.append(dict(
            operation='docx_append_paragraph', family='docx',
            request=f'Add one new paragraph with the exact text "{APPENDED_PARAGRAPH}" at the end of '
                    'the document body, after every existing paragraph. Change nothing else.',
            text=APPENDED_PARAGRAPH, paragraphCount=len(nodes)))
    return tasks, skipped


def plan_xlsx(parts):
    """Two tasks a workbook can state about itself: one typed write into an empty cell, one appended row."""
    tasks, skipped = [], []
    chosen = None
    for name, part, state in workbook_sheets(parts):
        if state != 'visible' or not name or part not in parts:
            continue
        if not re.fullmatch(r"[A-Za-z0-9 _.()-]{1,31}", name):
            continue
        grid = list(sheet_grid(parts[part]))
        if not grid:
            continue
        # The last used row is the greatest row the sheet carries, a row that holds no cell
        # of its own included, which is how a worksheet writer reads "after the last row".
        chosen = (name, part, grid, max(number for number, _ in sheet_rows(ET.fromstring(parts[part]))))
        break
    if chosen is None:
        skipped.append('no visible worksheet with a plainly addressable name and used cells was found')
        return tasks, skipped
    name, part, grid, last_row = chosen
    free_column = max(column for _, _, _, column in grid) + 1
    first_row = min(number for _, _, number, _ in grid)
    first_column = min(column for _, _, _, column in grid)
    if free_column > 16384:
        skipped.append(f'sheet "{name}" already uses the last worksheet column, so it states no empty cell here')
    else:
        cell = f'{column_name(free_column)}{first_row}'
        tasks.append(dict(
            operation='xlsx_set_cell', family='xlsx',
            request=f'On sheet "{name}", write the number {WRITTEN_NUMBER} into the empty cell {cell}. '
                    'Change nothing else.',
            sheet=name, sheetPart=part, cell=cell, value=WRITTEN_NUMBER, valueType='number'))
    if last_row + 1 > 1048576:
        skipped.append(f'sheet "{name}" already reaches the last worksheet row, so no row can be appended')
    else:
        tasks.append(dict(
            operation='xlsx_append_row', family='xlsx',
            request=f'On sheet "{name}", append one new row after the last used row, holding the text '
                    f'"{APPENDED_ROW}" in column {column_name(first_column)}. Change nothing else.',
            sheet=name, sheetPart=part, row=last_row + 1, column=column_name(first_column),
            columnNumber=first_column, text=APPENDED_ROW))
    return tasks, skipped


def plan_pptx(parts):
    """Two tasks a deck can state about itself: one word rewrite in a shape, one duplicated slide."""
    tasks, skipped = [], []
    slides = slide_parts(parts)
    if not slides:
        skipped.append('the presentation lists no slides')
        return tasks, skipped
    chosen = None
    for position, part in enumerate(slides):
        root = ET.fromstring(parts[part])
        holders = [n.text or '' for n in root.iter() if n.tag == f'{{{A}}}t']
        elsewhere = other_part_bytes(parts, part)
        taken = elsewhere + b'\n' + '\n'.join(holders).lower().encode()
        tree = root.find(f'{{{P}}}cSld/{{{P}}}spTree')
        if tree is None:
            continue
        for ordinal, shape in enumerate(tree.findall(f'{{{P}}}sp')):
            marker = shape.find(f'{{{P}}}nvSpPr/{{{P}}}cNvPr')
            body = shape.find(f'{{{P}}}txBody')
            if marker is None or body is None or not (marker.get('id') or '').isdigit():
                continue
            for index, paragraph in enumerate(body.findall(f'{{{A}}}p')):
                if not plain_slide_paragraph(paragraph):
                    continue
                for run in paragraph.findall(f'{{{A}}}r'):
                    for text_node in run.findall(f'{{{A}}}t'):
                        word = unique_word(text_node.text or '', holders, elsewhere)
                        replacement = free_replacement(taken, word or '') if word else None
                        if replacement is not None:
                            chosen = (position + 1, part, int(marker.get('id')), ordinal + 1, index,
                                      paragraph, word, replacement)
                            break
                    if chosen:
                        break
                if chosen:
                    break
            if chosen:
                break
        if chosen:
            break
    if chosen is None:
        skipped.append('no plain shape paragraph holds a word that occurs exactly once in the package')
    else:
        slide, part, shape_id, ordinal, index, paragraph, word, replacement = chosen
        tasks.append(dict(
            operation='pptx_replace_word', family='pptx',
            request=f'On slide {slide}, replace the word "{word}" with "{replacement}". Change nothing else.',
            word=word, replacement=replacement, slide=slide, slidePart=part, shapeId=shape_id,
            shapeOrdinal=ordinal, paragraphIndex=index, paragraphText=element_text(paragraph, A)))
    tasks.append(dict(
        operation='pptx_duplicate_slide', family='pptx',
        request='Duplicate slide 1 and place the copy immediately after it, as the new slide 2. '
                'Change nothing else.',
        slide=1, after=1, slidePart=slides[0], slideCount=len(slides)))
    return tasks, skipped


def plan_tasks(family, data):
    """Every task this file can state about itself, derived from its own bytes and nothing else."""
    if len(data) > CORPUS_MAX_BYTES:
        return [], [f'the file is over the {CORPUS_MAX_BYTES} byte comparison bound']
    if family not in CORPUS_FAMILIES:
        return [], [f'this comparison plans docx, xlsx and pptx tasks only, not {family}']
    try:
        parts = corpus_package(data)
        tasks, skipped = {'docx': plan_docx, 'xlsx': plan_xlsx, 'pptx': plan_pptx}[family](parts)
    except (zipfile.BadZipFile, ET.ParseError, OracleFailure, ValueError, KeyError, IndexError,
            TypeError, AttributeError, RecursionError) as exc:
        return [], [f'the package could not be planned for here: {type(exc).__name__}: {exc}']
    return tasks, skipped


# The relationship types whose target belongs to the one slide that references it. PowerPoint's
# own Duplicate Slide writes a new part for each of these and shares everything else, so a copy
# owns its charts, SmartArt, embeddings, VML, tags, comments and speaker notes while layouts,
# masters, themes and media stay shared. The same list the own-file recipe's duplicate oracle
# uses; the closure is re-derived from the PRODUCED package, never taken from a tool's own word.
OFFICE_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
PPTX_PRIVATE_REL_TYPES = {
    OFFICE_REL + name for name in (
        'notesSlide', 'chart', 'chartUserShapes', 'themeOverride', 'diagramData', 'diagramLayout',
        'diagramColors', 'diagramQuickStyle', 'oleObject', 'package', 'vmlDrawing', 'tags',
        'comments', 'control',
    )
} | {
    'http://schemas.microsoft.com/office/2007/relationships/diagramDrawing',
    'http://schemas.microsoft.com/office/2011/relationships/chartColorStyle',
    'http://schemas.microsoft.com/office/2011/relationships/chartStyle',
    'http://schemas.microsoft.com/office/2014/relationships/chartEx',
    'http://schemas.microsoft.com/office/2006/relationships/activeXControlBinary',
    'http://schemas.microsoft.com/office/2018/10/relationships/comments',
}
PART_IN_DETAIL = re.compile(r'\b(?:xl|word|ppt|customXml|docProps)/[A-Za-z0-9_./-]+\.(?:xml|rels|bin)\b')


def declared_parts(commands):
    """Every package member a tool itself NAMED as written, read from its own printed report.

    A tool that publishes a machine-readable change list declares what it wrote: UFO's edit
    receipt names the table part a row append grew (`Table1 (xl/tables/table1.xml) extended to
    A1:D3`) and the slide part a duplicate added. The rule is the same for every tool in the
    comparison, and a tool that declares nothing gets nothing admitted, which is the honest
    reading: an undeclared member outside the fixed write set is a foreign member.
    """
    named = set()
    for command in commands:
        try:
            payload = json.loads(command.get('stdout') or '')
        except ValueError:
            continue
        result = payload.get('result') if isinstance(payload, dict) else None
        applied = result.get('applied') if isinstance(result, dict) else None
        for row in applied if isinstance(applied, list) else []:
            if isinstance(row, dict) and isinstance(row.get('detail'), str):
                named |= set(PART_IN_DETAIL.findall(row['detail']))
    return named


def relationships_path(part):
    folder, _, name = part.rpartition('/')
    return f'{folder}/_rels/{name}.rels' if folder else f'_rels/{name}.rels'


def slide_owned_parts(parts, slides):
    """Every part [slides] own in [parts]: each slide, each part reached from it by a relationship
    type private to a slide, and the `_rels` part of each of those."""
    owned, pending = set(), [name for name in sorted(slides) if name in parts]
    seen = set(pending)
    while pending:
        part = pending.pop()
        owned.add(part)
        rels = relationships_path(part)
        if rels not in parts:
            continue
        owned.add(rels)
        try:
            relationships = list(ET.fromstring(parts[rels]))
        except ET.ParseError:
            continue
        for relationship in relationships:
            if relationship.get('TargetMode') == 'External' or relationship.get('Type') not in PPTX_PRIVATE_REL_TYPES:
                continue
            target = rel_target(rels, relationship.get('Target', ''))
            if target in parts and target not in seen:
                seen.add(target)
                pending.append(target)
    return owned


def corpus_allowed(task, before, after, declared=frozenset()):
    """The members this task may change, add or remove. The rule is identical for every tool.

    [declared] is what the tool itself named as written, and only two things turn on it, both
    because the fixed write set alone would call a correct edit wrong. Everything else stays as
    it was: a rewritten unrelated part, an added metadata part or a re-serialized untouched
    member is a foreign member whatever a tool says about it.
    """
    operation = task['operation']
    changed, added, removed, pattern = set(), set(), set(), None
    if operation in ('docx_replace_word', 'docx_append_paragraph'):
        changed = {'word/document.xml'}
    elif operation in ('xlsx_set_cell', 'xlsx_append_row'):
        # The recalculation policy and the formula-order cache belong to a cell write; the
        # shared string table belongs to a written string.
        changed = {task['sheetPart'], 'xl/workbook.xml'}
        removed = {'xl/calcChain.xml'}
        if operation == 'xlsx_append_row':
            changed.add('xl/sharedStrings.xml')
            added.add('xl/sharedStrings.xml')
            # A worksheet table that ends exactly at the last used row grows over the appended
            # rows, which is what Excel does and what the tool's own receipt names. Only a table
            # the tool NAMED, and only one the source already carries: a table part a tool is
            # silent about, or one that was added or removed, stays a foreign member.
            changed |= {name for name in declared
                        if name.startswith('xl/tables/') and name.endswith('.xml') and name in before}
        dropped = 'xl/calcChain.xml' in before and 'xl/calcChain.xml' not in after
        gained = 'xl/sharedStrings.xml' in added and 'xl/sharedStrings.xml' not in before
        if dropped or gained:
            changed |= {'[Content_Types].xml', 'xl/_rels/workbook.xml.rels'}
    elif operation == 'pptx_replace_word':
        changed = {task['slidePart']}
    elif operation == 'pptx_duplicate_slide':
        changed = {'[Content_Types].xml', 'ppt/presentation.xml', 'ppt/_rels/presentation.xml.rels'}
        pattern = r'ppt/(slides|notesSlides)/(_rels/)?(slide|notesSlide)\d+\.xml(\.rels)?'
        # A duplicate may add the copy's own parts and nothing else. PowerPoint's own Duplicate
        # Slide copies a slide's private closure (its charts and their rels, embeddings, activeX,
        # VML drawings, diagram parts, tags, comments and speaker notes) into new parts and shares
        # everything else. The closure is walked in the PRODUCED package from the slide parts the
        # tool named, so a part a tool copied without wiring it to the copy, or added for no slide
        # at all, is still a foreign member. A tool that names no new slide may add nothing.
        copies = {name for name in declared if re.fullmatch(r'ppt/slides/slide\d+\.xml', name)}
        added = {name for name in slide_owned_parts(after, copies) if name not in before}
    return changed, added, removed, pattern


def cells_of(parts, part):
    strings = [text_of(n, S) for n in ET.fromstring(parts.get('xl/sharedStrings.xml', f'<sst xmlns="{S}"/>'.encode())).findall('s:si', NS)]
    result = {}
    for reference, c, _, _ in sheet_grid(parts[part]):
        kind, value = c.get('t', 'n'), c.findtext('s:v', default='', namespaces=NS)
        formula = c.findtext('s:f', namespaces=NS)
        if formula is not None:
            result[reference] = ('formula', formula)
        elif kind in ('s', 'inlineStr', 'str'):
            if kind == 'inlineStr':
                result[reference] = ('string', text_of(c, S))
            elif kind == 'str':
                result[reference] = ('string', value)
            else:
                # A shared cell must carry an index into the table. Anything else is reported
                # as broken rather than guessed at, because Excel would read an index here.
                index = int(value) if value.strip().isdigit() else -1
                result[reference] = (('string', strings[index]) if 0 <= index < len(strings)
                                     else ('broken-shared-string', value))
        elif kind == 'b':
            result[reference] = ('boolean', value in ('1', 'true', 'TRUE'))
        elif value != '':
            result[reference] = ('number', value)
    return result


def same_cell(first, second):
    """Cell equality that forgives only what a written cell legitimately changes.

    A number is compared as a number, so 12.5 and 12.50 agree. A shared-formula dependent
    stores no formula text of its own, so a tool that materializes the group's formula is not
    counted as a content difference; that rewrite still shows on the preservation axis.
    """
    if first == second:
        return True
    if not (isinstance(first, tuple) and isinstance(second, tuple)) or first[0] != second[0]:
        return False
    if first[0] == 'formula':
        return '' in (first[1], second[1])
    if first[0] != 'number':
        return False
    try:
        return float(first[1]) == float(second[1])
    except (TypeError, ValueError):
        return False


def first_difference(expected, actual, label):
    for index in range(max(len(expected), len(actual))):
        one = expected[index] if index < len(expected) else '<missing>'
        other = actual[index] if index < len(actual) else '<missing>'
        if one != other:
            return (f'{label} {index} differs: expected {str(one)[:80]!r}, found {str(other)[:80]!r} '
                    f'({len(expected)} expected, {len(actual)} found)')
    return ''


def slide_texts(parts):
    return [element_text(ET.fromstring(parts[name]), A) for name in slide_parts(parts)]


def docx_texts(parts):
    return [element_text(node, W) for node, _ in docx_nodes(docx_body(parts))]


def corpus_content(task, before, after):
    """Did exactly the requested change land, read back by this runner's own parsers."""
    operation = task['operation']
    if operation == 'docx_replace_word':
        expected = docx_texts(before)
        index = task['paragraphIndex']
        require(index < len(expected), 'the source paragraph is gone')
        expected[index] = expected[index].replace(task['word'], task['replacement'], 1)
        actual = docx_texts(after)
        require(actual == expected, first_difference(expected, actual, 'body paragraph'))
    elif operation == 'docx_append_paragraph':
        expected = docx_texts(before) + [task['text']]
        actual = docx_texts(after)
        require(actual == expected, first_difference(expected, actual, 'body paragraph'))
        body = docx_body(after)
        direct = [child for child in body if child.tag == f'{{{W}}}p']
        require(direct and element_text(direct[-1], W) == task['text'],
                'the new paragraph is not the last paragraph of the body itself')
    elif operation in ('xlsx_set_cell', 'xlsx_append_row'):
        part = task['sheetPart']
        expected = cells_of(before, part)
        if operation == 'xlsx_set_cell':
            expected[task['cell']] = ('number', task['value'])
        else:
            expected[f'{task["column"]}{task["row"]}'] = ('string', task['text'])
        actual = cells_of(after, part)
        wrong = sorted(ref for ref in expected.keys() | actual.keys()
                       if not same_cell(expected.get(ref), actual.get(ref)))
        require(not wrong, f'cells differ at {wrong[:6]}: expected '
                           f'{[expected.get(r) for r in wrong[:3]]}, found {[actual.get(r) for r in wrong[:3]]}')
    elif operation == 'pptx_replace_word':
        expected = slide_texts(before)
        index = task['slide'] - 1
        require(index < len(expected), 'the source slide is gone')
        expected[index] = expected[index].replace(task['word'], task['replacement'], 1)
        actual = slide_texts(after)
        require(actual == expected, first_difference(expected, actual, 'slide'))
    elif operation == 'pptx_duplicate_slide':
        original = slide_texts(before)
        expected = original[:task['after']] + [original[task['slide'] - 1]] + original[task['after']:]
        actual = slide_texts(after)
        require(actual == expected, first_difference(expected, actual, 'slide'))
    else:
        raise OracleFailure(f'no content oracle for {operation}')
    return True


def sole_text_node(root, namespace, word):
    found = [node for node in root.iter(f'{{{namespace}}}t') if word in (node.text or '')]
    require(len(found) == 1, 'the planned word is no longer unique in the source part')
    return found[0]


def normalized_sheet(data, drop_cells=(), drop_rows=()):
    """The worksheet without the task's own cells or rows, its used-range note or refreshed caches.

    A cell write may legitimately change the used range, the row's `spans` hint and the cached
    value of a formula that reads the written cell. Nothing else in the grid may move, and the
    same allowance is given to every tool.
    """
    root = ET.fromstring(data)
    grid = root.find('s:sheetData', NS)
    for number, row in list(sheet_rows(root)):
        if str(number) in drop_rows:
            grid.remove(row)
            continue
        row.attrib.pop('spans', None)
        column = 0
        for cell in list(row.findall('s:c', NS)):
            reference = cell.get('r') or ''
            if re.fullmatch(r'[A-Z]+[0-9]+', reference):
                column = column_number(re.match(r'[A-Z]+', reference).group(0))
            else:
                column += 1
                reference = f'{column_name(column)}{number}'
            if reference in drop_cells:
                row.remove(cell)
            elif cell.find('s:f', NS) is not None:
                for cached in cell.findall('s:v', NS):
                    cell.remove(cached)
    for node in root.findall('s:dimension', NS):
        root.remove(node)
    return root


def normalized_workbook(data):
    root = ET.fromstring(data)
    for node in root.findall('s:calcPr', NS):
        root.remove(node)
    return root


def relationship_set(data):
    return {(r.get('Id'), r.get('Type'), r.get('Target'), r.get('TargetMode')) for r in ET.fromstring(data)}


def content_type_set(data):
    return {(node.tag, tuple(sorted(node.attrib.items()))) for node in ET.fromstring(data)}


def corpus_residual(task, before, after):
    """Everything in the edited members except the requested change must still be the same document."""
    operation = task['operation']
    if operation in ('docx_replace_word', 'pptx_replace_word'):
        part = 'word/document.xml' if operation == 'docx_replace_word' else task['slidePart']
        namespace = W if operation == 'docx_replace_word' else A
        expected = ET.fromstring(before[part])
        node = sole_text_node(expected, namespace, task['word'])
        node.text = (node.text or '').replace(task['word'], task['replacement'], 1)
        require(xml_shape(ET.fromstring(after[part])) == xml_shape(expected),
                'the edited part is not the source with only the requested characters changed')
    elif operation == 'docx_append_paragraph':
        root = ET.fromstring(after['word/document.xml'])
        body = root.find(f'{{{W}}}body')
        direct = [child for child in body if child.tag == f'{{{W}}}p']
        require(direct, 'the produced body holds no paragraph')
        body.remove(direct[-1])
        require(xml_shape(root) == xml_shape(ET.fromstring(before['word/document.xml'])),
                'the rest of the document part was rewritten by the append')
    elif operation in ('xlsx_set_cell', 'xlsx_append_row'):
        part = task['sheetPart']
        if operation == 'xlsx_set_cell':
            number = re.search(r'\d+', task['cell']).group(0)
            rows = {row.get('r') for row in ET.fromstring(before[part]).findall('s:sheetData/s:row', NS)}
            drop_cells, drop_rows = ({task['cell']}, set()) if number in rows else (set(), {number})
        else:
            drop_cells, drop_rows = set(), {str(task['row'])}
        first = normalized_sheet(before[part], drop_cells, drop_rows)
        second = normalized_sheet(after[part], drop_cells, drop_rows)
        require(xml_shape(first) == xml_shape(second), 'the worksheet outside the written cell changed')
        require(xml_shape(normalized_workbook(before['xl/workbook.xml'])) ==
                xml_shape(normalized_workbook(after.get('xl/workbook.xml', before['xl/workbook.xml']))),
                'workbook content other than the recalculation policy changed')
        if before.get('xl/sharedStrings.xml') != after.get('xl/sharedStrings.xml') and 'xl/sharedStrings.xml' in before:
            old = ET.fromstring(before['xl/sharedStrings.xml']).findall('s:si', NS)
            new = ET.fromstring(after['xl/sharedStrings.xml']).findall('s:si', NS)
            require([xml_shape(n) for n in new[:len(old)]] == [xml_shape(n) for n in old],
                    'existing shared strings were rewritten or reordered')
    elif operation == 'pptx_duplicate_slide':
        first = ET.fromstring(before['ppt/presentation.xml'])
        second = ET.fromstring(after['ppt/presentation.xml'])
        listed = second.find('p:sldIdLst', NS)
        require(listed is not None and len(listed) == task['slideCount'] + 1, 'the slide list did not grow by one')
        listed.remove(list(listed)[task['after']])
        require(xml_shape(first) == xml_shape(second), 'the presentation part changed beyond the new slide')
        gained = relationship_set(after['ppt/_rels/presentation.xml.rels']) - relationship_set(before['ppt/_rels/presentation.xml.rels'])
        require(relationship_set(before['ppt/_rels/presentation.xml.rels']) <=
                relationship_set(after['ppt/_rels/presentation.xml.rels']),
                'an existing presentation relationship was rewritten or removed')
        require(all(r[1] and r[1].endswith('/slide') for r in gained), 'a relationship other than the new slide was added')
        require(content_type_set(before['[Content_Types].xml']) <= content_type_set(after['[Content_Types].xml']),
                'an existing content type entry was rewritten or removed')
    return True


def corpus_opens(data):
    """A cheap third check with this runner's own parsers: the package still reads as OPC."""
    parts = corpus_package(data)
    require('[Content_Types].xml' in parts, 'the package carries no [Content_Types].xml')
    relationships_valid(parts)
    return True


def verdict(check, *args):
    try:
        check(*args)
    except (OracleFailure, ET.ParseError, zipfile.BadZipFile, ValueError, KeyError, IndexError,
            TypeError, AttributeError, RecursionError) as exc:
        reason = str(exc) if isinstance(exc, OracleFailure) else f'{type(exc).__name__}: {exc}'
        kind = getattr(exc, 'checks', {}).get('kind', 'unreadable')
        return {'pass': False, 'reason': reason[:400], 'kind': kind}
    return {'pass': True, 'reason': None, 'kind': None}


def corpus_preservation(task, before, after, declared=frozenset(), stamp_aware=True):
    """Untouched members byte-identical and the edited members unchanged beyond the request.

    With [stamp_aware], the provenance and save stamps a tool adds are set aside first (see
    tools/corpus/package_semantics.py), exactly those and nothing else; without it this is the
    strict byte rule, which every record keeps beside the stamp-aware verdict.
    """
    if stamp_aware:
        before, after, _ = package_semantics.stamp_filter(before, after)
    changed, added, removed, pattern = corpus_allowed(task, before, after, declared)
    problems = []
    # An ADDED member is judged by the permitted set alone. A pattern over part names would admit
    # any `slideN.xml`, including one no slide of the produced deck owns.
    for name in sorted(after.keys() - before.keys()):
        if name not in added:
            problems.append(f'added a member the change does not need: {name}')
    for name in sorted(before.keys() - after.keys()):
        if name not in removed:
            problems.append(f'removed a member the change does not need: {name}')
    for name in sorted(before.keys() & after.keys()):
        if before[name] != after[name] and name not in changed:
            problems.append(f'rewrote an untouched member: {name}')
    if problems:
        # Two very different failures share this axis: touching a member the change never needed,
        # and rewriting the edited member beyond the characters that were asked for. The report
        # keeps them apart, because they mean different things to a reader.
        raise OracleFailure('; '.join(problems[:8]) + (f' (+{len(problems) - 8} more)' if len(problems) > 8 else ''),
                            {'kind': 'foreign_member'})
    try:
        corpus_residual(task, before, after)
    except OracleFailure as exc:
        raise OracleFailure(str(exc), {'kind': 'edited_part'}) from exc
    return True


def corpus_grade(task, before, after, declared=frozenset()):
    """What differs outside the write set, graded: preserved, rewritten (same content) or changed.

    The write set is the one `corpus_preservation` uses, so a `preserved` grade and a passing
    stamp-aware byte rule agree on every member outside the edited ones.
    """
    try:
        changed, added, removed, _ = corpus_allowed(task, before, after, declared)
        return package_semantics.grade(before, after, changed, added, removed, task['family'])
    except (OracleFailure, ET.ParseError, ValueError, KeyError, IndexError, TypeError, AttributeError,
            RecursionError) as exc:
        return {'verdict': None, 'reason': f'the grade could not be computed: {type(exc).__name__}: {exc}'[:400]}


def corpus_oracles(task, original, produced, source_opens, declared=frozenset()):
    """Three independent verdicts on the bytes alone: opens, requested content, preservation.

    Preservation carries two more readings beside its pass: `strict`, the byte rule with no
    stamp filter, and `grade`, the three-way verdict on everything outside the write set.
    """
    checks = {'opens': verdict(corpus_opens, produced) if source_opens else
              {'pass': None, 'reason': 'the source itself does not pass this check, so the output is not judged by it'}}
    try:
        before, after = corpus_package(original), corpus_package(produced)
    except (zipfile.BadZipFile, ValueError) as exc:
        reason = f'the produced file is not a readable package: {exc}'
        checks['content'] = {'pass': False, 'reason': reason}
        checks['preservation'] = {'pass': False, 'reason': reason}
        return checks
    checks['content'] = verdict(corpus_content, task, before, after)
    checks['preservation'] = verdict(corpus_preservation, task, before, after, declared)
    checks['preservation']['strict'] = verdict(corpus_preservation, task, before, after, declared, False)
    checks['preservation']['grade'] = corpus_grade(task, before, after, declared)
    return checks


DOCX_REPLACE_SCRIPT = '''"""python-docx: rewrite the one run that holds the word."""
import json, sys
import docx

source, output, task = sys.argv[1], sys.argv[2], json.loads(sys.argv[3])
document = docx.Document(source)
for paragraph in document.paragraphs:
    for run in paragraph.runs:
        if task["word"] in run.text:
            run.text = run.text.replace(task["word"], task["replacement"], 1)
            document.save(output)
            raise SystemExit(0)
raise SystemExit("python-docx found no body run holding the word")
'''
DOCX_APPEND_SCRIPT = '''"""python-docx: append one paragraph at the end of the body."""
import json, sys
import docx

source, output, task = sys.argv[1], sys.argv[2], json.loads(sys.argv[3])
document = docx.Document(source)
document.add_paragraph(task["text"])
document.save(output)
'''
XLSX_CELL_SCRIPT = '''"""openpyxl: write one typed number into an empty cell."""
import json, sys
from openpyxl import load_workbook

source, output, task = sys.argv[1], sys.argv[2], json.loads(sys.argv[3])
book = load_workbook(source)
book[task["sheet"]][task["cell"]] = float(task["value"])
book.save(output)
'''
XLSX_ROW_SCRIPT = '''"""openpyxl: write one text cell into the first row after the used range."""
import json, sys
from openpyxl import load_workbook

source, output, task = sys.argv[1], sys.argv[2], json.loads(sys.argv[3])
book = load_workbook(source)
book[task["sheet"]].cell(row=task["row"], column=task["columnNumber"], value=task["text"])
book.save(output)
'''
PPTX_REPLACE_SCRIPT = '''"""python-pptx: rewrite the one run of the named shape that holds the word."""
import json, sys
from pptx import Presentation

source, output, task = sys.argv[1], sys.argv[2], json.loads(sys.argv[3])
deck = Presentation(source)
slide = deck.slides[task["slide"] - 1]
for shape in slide.shapes:
    if shape.shape_id != task["shapeId"] or not shape.has_text_frame:
        continue
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if task["word"] in run.text:
                run.text = run.text.replace(task["word"], task["replacement"], 1)
                deck.save(output)
                raise SystemExit(0)
raise SystemExit("python-pptx found no run of the named shape holding the word")
'''
LIBRARY_SCRIPTS = {
    ('python-docx', 'docx_replace_word'): DOCX_REPLACE_SCRIPT,
    ('python-docx', 'docx_append_paragraph'): DOCX_APPEND_SCRIPT,
    ('openpyxl', 'xlsx_set_cell'): XLSX_CELL_SCRIPT,
    ('openpyxl', 'xlsx_append_row'): XLSX_ROW_SCRIPT,
    ('python-pptx', 'pptx_replace_word'): PPTX_REPLACE_SCRIPT,
}
LIBRARY_GAPS = {
    ('python-pptx', 'pptx_duplicate_slide'):
        'python-pptx 1.0.2 exposes no slide copy; its documentation and issue tracker direct callers '
        'to rebuild the slide by hand, which is a different edit from duplicating one.',
}


def corpus_ufo_adapter(task, runner, source, output, variant=0):
    """UFO reads the file's own coordinates and hash first, then writes one guarded batch."""
    operation = task['operation']

    def structure(arguments):
        return ufo_payload(runner.json(['text', '--format', 'structure', *arguments, str(source)]))

    if operation == 'docx_replace_word':
        page = structure(['--offset', '0', '--limit', '200'])
        digest, total, offset = page['source']['sha256'], page.get('totalParagraphs', 0), 0
        rows = page.get('paragraphs') or []
        row = next((r for r in rows if task['word'] in (r.get('text') or '')), None)
        while row is None and offset + len(rows) < total and offset < 800:
            offset += len(rows)
            rows = structure(['--offset', str(offset), '--limit', '200']).get('paragraphs') or []
            if not rows:
                break
            row = next((r for r in rows if task['word'] in (r.get('text') or '')), None)
        if row is None:
            raise CommandError('the structured read never reported a paragraph holding the word')
        text = row['text']
        arguments = ['edit', 'paragraphs', '--expect-sha256', digest,
                     '--expect', f'{row["index"]}={text}',
                     '--set', f'{row["index"]}={text.replace(task["word"], task["replacement"], 1)}']
    elif operation == 'docx_append_paragraph':
        page = structure(['--offset', '0', '--limit', '1'])
        total = page.get('totalParagraphs', 0)
        if not total:
            raise CommandError('the structured read reported no paragraph to append after')
        arguments = ['edit', 'paragraphs', '--expect-sha256', page['source']['sha256'],
                     '--insert-after', f'{total - 1}={task["text"]}']
    elif operation == 'xlsx_set_cell':
        page = structure(['--sheet', task['sheet'], '--range', f'{task["cell"]}:{task["cell"]}'])
        found = next((c for c in page['selection']['cells'] if c['ref'] == task['cell']), None)
        if found is None:
            raise CommandError('the structured read did not report the target cell')
        preimage = 'empty' if found['type'] == 'empty' else f'{found["type"]}:{found["value"]}'
        arguments = ['edit', 'cells', '--expect-sha256', page['source']['sha256'],
                     '--expect', f'{task["sheet"]}!{task["cell"]}={preimage}',
                     '--set', f'{task["sheet"]}!{task["cell"]}=number:{task["value"]}']
    elif operation == 'xlsx_append_row':
        page = structure(['--sheet', task['sheet'], '--range', 'A1:A1'])
        arguments = ['edit', 'rows', '--expect-sha256', page['source']['sha256'], '--sheet', task['sheet'],
                     '--append-row', json.dumps({task['column']: f'string:{task["text"]}'})]
    elif operation == 'pptx_replace_word':
        page = structure(['--slide', str(task['slide'])])
        rows = [p for shape in page['selection']['shapes'] for p in (shape.get('paragraphs') or [])]
        row = next((p for p in rows if task['word'] in (p.get('text') or '')), None)
        if row is None:
            raise CommandError('the slide read never reported a paragraph holding the word')
        text = row['text']
        arguments = ['edit', 'paragraphs', '--expect-sha256', page['source']['sha256'],
                     '--slide', str(task['slide']), '--expect', f'{row["index"]}={text}',
                     '--set', f'{row["index"]}={text.replace(task["word"], task["replacement"], 1)}']
    else:
        page = structure(['--slide', str(task['slide'])])
        arguments = ['edit', 'slides', '--expect-sha256', page['source']['sha256'],
                     '--duplicate', str(task['slide']), '--after', str(task['after'])]
    record = runner.run([*arguments, '-o', str(output), str(source)])
    if record['returncode']:
        raise CommandError(ufo_refusal(record))


def ufo_refusal(record):
    try:
        receipt = json.loads(record['stdout'])
    except ValueError:
        return f'UFO exited {record["returncode"]} without a receipt'
    return f'{receipt.get("code") or "failed"}: {receipt.get("message") or receipt.get("status")}'


def corpus_officecli_adapter(task, runner, source, output, variant=0):
    """OfficeCLI edits a staging copy in place, through its documented path addressing."""
    operation = task['operation']
    shutil.copyfile(source, output)
    name = str(output)
    if operation == 'docx_replace_word':
        path = task['paragraphPath'] if variant == 0 else '/body'
        arguments = ['set', name, path, '--find', task['word'], '--replace', task['replacement']]
    elif operation == 'docx_append_paragraph':
        arguments = ['add', name, '/body', '--type', 'paragraph', '--prop', 'text=' + task['text']]
    elif operation == 'xlsx_set_cell':
        arguments = ['set', name, f'/{task["sheet"]}/{task["cell"]}',
                     '--prop', 'value=' + task['value'], '--prop', 'type=number']
    elif operation == 'xlsx_append_row':
        arguments = ['set', name, f'/{task["sheet"]}/{task["column"]}{task["row"]}',
                     '--prop', 'value=' + task['text'], '--prop', 'type=string']
    elif operation == 'pptx_replace_word':
        path = f'/slide[{task["slide"]}]/shape[{task["shapeOrdinal"]}]' if variant == 0 else f'/slide[{task["slide"]}]'
        arguments = ['set', name, path, '--find', task['word'], '--replace', task['replacement']]
    else:
        arguments = ['add', name, '/', '--type', 'slide', '--from', f'/slide[{task["slide"]}]',
                     '--index', str(task['after'])]
    record = runner.run([*arguments, '--json'])
    if record['returncode']:
        raise CommandError(f'OfficeCLI exited {record["returncode"]}: {record["stderr"][:200] or record["stdout"][:200]}')


def corpus_docxcli_adapter(task, runner, source, output, variant=0):
    """docx-cli replaces the first exact match and appends at the end of the document."""
    operation = task['operation']
    if task['family'] != 'docx':
        raise Unsupported('docx-cli 0.25.0 reads and writes DOCX only.')
    if operation == 'docx_replace_word':
        arguments = ['replace', str(source), task['word'], task['replacement'], '--exact']
    else:
        arguments = ['insert', str(source), '--at-end', '--text', task['text']]
    record = runner.run([*arguments, '-o', str(output)])
    if record['returncode']:
        raise CommandError(f'docx-cli exited {record["returncode"]}: {record["stderr"][:200] or record["stdout"][:200]}')


def corpus_genoffice_adapter(task, runner, source, output, variant=0):
    """GenOffice applies one batch of its documented ops to the source and publishes with --out.

    A word replacement is scoped first (the one Word block or the one slide holding the word) and,
    as the declared second attempt, run over the whole document, as OfficeCLI's is. A cell write
    carries an empty expectation, which GenOffice checks before it writes.
    """
    operation = task['operation']
    if operation == 'docx_replace_word':
        replacement = {'op': 'findReplace', 'find': task['word'], 'replace': task['replacement'], 'matchCase': True}
        if variant == 0:
            blocks = genoffice_result(runner, ['docs', 'read', str(source), '--full'])['items']
            holder = next((b['index'] for b in blocks if task['word'] in (b.get('text') or '')), None)
            if holder is None:
                raise CommandError('the GenOffice read never reported a block holding the word')
            replacement['target'] = {'blockIndexes': [holder]}
        genoffice_apply(runner, 'docs', source, output, [replacement])
    elif operation == 'docx_append_paragraph':
        # Without an index, insert_content appends at the end of the document.
        genoffice_apply(runner, 'docs', source, output, [{'op': 'insert_content', 'html': f'<p>{task["text"]}</p>'}])
    elif operation == 'xlsx_set_cell':
        genoffice_apply(runner, 'sheet', source, output, [{'op': 'set_cell', 'sheet': task['sheet'], 'address': task['cell'],
                                                           'value': float(task['value']), 'expectedValue': None}])
    elif operation == 'xlsx_append_row':
        genoffice_apply(runner, 'sheet', source, output, [{'op': 'set_cell', 'sheet': task['sheet'],
                                                           'address': f'{task["column"]}{task["row"]}',
                                                           'value': task['text'], 'expectedValue': None}])
    elif operation == 'pptx_replace_word':
        replacement = {'op': 'findReplace', 'find': task['word'], 'replace': task['replacement'], 'matchCase': True}
        if variant == 0:
            replacement['slideIndex'] = task['slide'] - 1
        genoffice_apply(runner, 'slides', source, output, [replacement])
    else:
        # The copy lands immediately after the source, the position the task names.
        genoffice_apply(runner, 'slides', source, output, [{'op': 'duplicateSlide', 'target': {'slide': task['slide'] - 1}}])


def corpus_library_adapter(tool):
    """One short script per task, the route an agent takes when it reaches for the library."""
    def adapter(task, runner, source, output, variant=0):
        if LIBRARY_FAMILY[tool] != task['family']:
            raise Unsupported(f'{tool} reads and writes {LIBRARY_FAMILY[tool].upper()} only.')
        gap = LIBRARY_GAPS.get((tool, task['operation']))
        if gap:
            raise Unsupported(gap)
        script = LIBRARY_SCRIPTS.get((tool, task['operation']))
        if script is None:
            raise Unsupported(f'{tool} has no documented route for {task["operation"]}.')
        path = runner.cwd / 'edit.py'
        path.write_text(script, encoding='utf-8')
        payload = {k: v for k, v in task.items() if k not in ('request', 'id', 'file')}
        record = runner.run([str(path), str(source), str(output), json.dumps(payload, sort_keys=True)])
        record['script'] = script
        if record['returncode']:
            raise CommandError(f'{tool} exited {record["returncode"]}: '
                               f'{(record["stderr"] or record["stdout"]).strip().splitlines()[-1][:200] if (record["stderr"] or record["stdout"]).strip() else "no message"}')
    return adapter


CORPUS_ADAPTERS = {
    'ufo': corpus_ufo_adapter,
    'officecli': corpus_officecli_adapter,
    'genoffice': corpus_genoffice_adapter,
    'docx-cli': corpus_docxcli_adapter,
    'python-docx': corpus_library_adapter('python-docx'),
    'openpyxl': corpus_library_adapter('openpyxl'),
    'python-pptx': corpus_library_adapter('python-pptx'),
}
CORPUS_VARIANTS = {('officecli', 'docx_replace_word'): 2, ('officecli', 'pptx_replace_word'): 2,
                   ('genoffice', 'docx_replace_word'): 2, ('genoffice', 'pptx_replace_word'): 2}


def library_requirements():
    return ''.join(f'{name}=={version} \\\n    --hash=sha256:{digest}\n' for name, version, digest in LIBRARY_PINS)


def library_probe():
    names = json.dumps([name for name, _, _ in LIBRARY_PINS])
    return ('import json, sys\nfrom importlib import metadata\n'
            f'print(json.dumps(dict([(n, metadata.version(n)) for n in {names}], python=sys.version.split()[0])))\n')


def provision_libraries(venv, install=False):
    """A private, hash-pinned environment under build/. Never the system interpreter."""
    python = venv / 'bin' / 'python'
    if not python.is_file():
        if not install:
            raise ValueError(f'the pinned library environment is missing: {venv}; create it once with '
                             'compare --corpus MANIFEST --setup-libraries')
        venv.parent.mkdir(parents=True, exist_ok=True)
        requirements = venv.parent / 'library-pins.txt'
        requirements.write_text(library_requirements(), encoding='utf-8')
        for command in ([sys.executable, '-m', 'venv', str(venv)],
                        [str(python), '-m', 'pip', 'install', '--require-hashes', '--no-deps',
                         '--disable-pip-version-check', '-r', str(requirements)]):
            done = subprocess.run(command, capture_output=True, text=True, timeout=1800)
            if done.returncode:
                raise ValueError(f'library provisioning failed: {" ".join(command[-3:])}: {done.stderr[-400:]}')
    probe = subprocess.run([str(python), '-c', library_probe()], capture_output=True, text=True, timeout=180)
    if probe.returncode:
        raise ValueError(f'the pinned library environment could not be read: {probe.stderr[-300:]}')
    installed = json.loads(probe.stdout)
    versions = {}
    for name, version, digest in LIBRARY_PINS:
        if installed.get(name) != version:
            raise ValueError(f'{name} is {installed.get(name)!r} in {venv}, not the pinned {version}')
        versions[name] = dict(version=version, sha256=digest)
    return {tool: dict(path=str(python), version=f'{tool} {versions[tool]["version"]}',
                       sha256=versions[tool]['sha256'], license=LIBRARY_LICENSE[tool],
                       interpreter=installed['python'], environment=str(venv),
                       pins={name: value['version'] for name, value in versions.items()})
            for tool in LIBRARY_FAMILY}


def corpus_files(manifest, root, families, limit=None):
    """Manifest rows that exist, hash to real bytes and fall inside the size bound."""
    rows, skipped = [], []
    with manifest.open(newline='', encoding='utf-8') as handle:
        for row in csv.DictReader(handle, delimiter='\t'):
            family = (row.get('family') or '').strip()
            if family not in families:
                # Recorded, never dropped in silence: a reader must see what the run left out.
                skipped.append(dict(file=row['id'], family=family, reasons=[
                    f'no tool in this comparison offers a {family} route, and this runner has no '
                    f'independent {family} oracle for real files'
                    if family not in CORPUS_FAMILIES else
                    f'--corpus-families selected {", ".join(families)}']))
                continue
            path = root / row['local_path']
            if not path.is_file():
                skipped.append(dict(file=row['id'], family=family, reasons=['the manifest row has no downloaded file here']))
                continue
            size = path.stat().st_size
            if size > CORPUS_MAX_BYTES:
                skipped.append(dict(file=row['id'], family=family,
                                    reasons=[f'the file is {size} bytes, over the {CORPUS_MAX_BYTES} byte bound']))
                continue
            rows.append(dict(id=row['id'], family=family, path=path, relative=row['local_path'],
                             producer=row.get('producer') or '', lineage=row.get('lineage') or '',
                             holdout=(row.get('holdout') or 'no').strip() == 'yes', bytes=size))
            if limit and len(rows) >= limit:
                break
    return rows, skipped


def corpus_plan(rows):
    """Every task the corpus states about itself, plus a reason for every file that yields none."""
    planned, skipped = [], []
    for row in rows:
        data = row['path'].read_bytes()
        tasks, reasons = plan_tasks(row['family'], data)
        opens = verdict(corpus_opens, data)['pass'] if tasks else False
        for task in tasks:
            planned.append(dict(task, id=f'{row["id"]}::{task["operation"]}', file=row['id'],
                                sourcePath=row['relative'], sourceSha256=sha(data), sourceOpens=opens))
        if reasons:
            skipped.append(dict(file=row['id'], family=row['family'], reasons=reasons))
    return planned, skipped


def execute_corpus_task(task, tool, identity, data, prefix, outputs, timeout):
    """One tool, one task, one fresh copy, judged only on the bytes it published."""
    start = time.perf_counter()
    record = dict(schemaVersion=CORPUS_VERSION, id=task['id'], file=task['file'], family=task['family'],
                  task=task['operation'], tool=tool, result='error', reason=None, checks={},
                  output=None, commands=[], commandCount=0, stdoutBytes=0, seconds=0.0,
                  sourceSha256=task['sourceSha256'],
                  # The coordinates the task states about itself, so a later step can tell
                  # the cells this edit wrote from the cells it only carried over.
                  sheet=task.get('sheet'), cell=task.get('cell'), row=task.get('row'))
    variants = CORPUS_VARIANTS.get((tool, task['operation']), 1)
    with tempfile.TemporaryDirectory(prefix='corpus-') as temporary:
        cwd = Path(temporary)
        source = cwd / ('input.' + task['family'])
        output = cwd / ('output.' + task['family'])
        source.write_bytes(data)
        runner = Runner(identity['path'], cwd, prefix, timeout=timeout)
        produced = None
        try:
            for variant in range(variants):
                # Each declared attempt starts from nothing, so the published file and the
                # recorded verdict always come from the same attempt.
                produced = None
                if output.exists():
                    output.unlink()
                try:
                    CORPUS_ADAPTERS[tool](task, runner, source, output, variant)
                except CommandError as exc:
                    reason = str(exc)
                    if source.read_bytes() != data:
                        record.update(result='wrong', reason='the tool changed its own source copy')
                    elif reason in ('timeout', 'capture_limit'):
                        record.update(result='error', reason=f'the command hit the {reason} bound')
                    elif output.exists() and output.read_bytes() != data:
                        record.update(result='error',
                                      reason=f'the command failed after changing its own copy: {reason}')
                    else:
                        # Nothing of the tool's own published and the input untouched: a documented
                        # refusal and a failure that wrote nothing look the same from outside, so
                        # they share a bucket rather than being guessed apart. A tool that edits a
                        # staging copy in place has published nothing while that copy is unchanged.
                        record.update(result='declined', reason=reason)
                    continue
                if source.read_bytes() != data:
                    record.update(result='wrong', reason='the tool changed its own source copy')
                    continue
                if not output.exists():
                    record.update(result='error', reason='the command reported success and published no file')
                    continue
                produced = output.read_bytes()
                checks = corpus_oracles(task, data, produced, task['sourceOpens'],
                                        declared_parts(runner.commands))
                record['checks'] = checks
                record['result'] = 'pass' if checks['content']['pass'] else 'wrong'
                record['reason'] = None if checks['content']['pass'] else checks['content']['reason']
                if record['result'] == 'pass':
                    break
        except Unsupported as exc:
            record.update(result='unsupported', reason=str(exc))
        except (OSError, ValueError, KeyError, IndexError, TypeError, StopIteration, RecursionError) as exc:
            record.update(result='error', reason=f'adapter or runtime: {type(exc).__name__}: {exc}')
        if produced is not None:
            folder = outputs / tool / task['file']
            folder.mkdir(parents=True, exist_ok=True)
            destination = folder / f'{task["operation"]}.{task["family"]}'
            destination.write_bytes(produced)
            record['output'] = dict(path=str(destination.relative_to(outputs.parent)),
                                    sha256=sha(produced), bytes=len(produced))
        record.update(commands=[corpus_command(c) for c in runner.commands], commandCount=len(runner.commands),
                      stdoutBytes=sum(c['stdout_bytes'] for c in runner.commands),
                      seconds=time.perf_counter() - start)
    return record


def corpus_command(command):
    return dict(argv=command['argv'], returncode=command.get('returncode'), seconds=command.get('wall_seconds'),
                stdoutBytes=command.get('stdout_bytes', 0), stderrBytes=command.get('stderr_bytes', 0),
                stdout=(command.get('stdout') or '')[:2000], stderr=(command.get('stderr') or '')[:2000],
                script=command.get('script'), error=command.get('error'))


def rate(numerator, denominator):
    return round(numerator / denominator, 4) if denominator else None


def corpus_summarize(records, tools):
    """Per family and overall: what each tool was asked, what it did and what it left alone."""
    summary = {}
    for scope in [*sorted({r['family'] for r in records}), 'all']:
        rows_in_scope = [r for r in records if scope == 'all' or r['family'] == scope]
        table = {}
        for tool in tools:
            rows = [r for r in rows_in_scope if r['tool'] == tool]
            counts = Counter(r['result'] for r in rows)
            attempted = len(rows) - counts['unsupported']
            produced = [r for r in rows if r['result'] in ('pass', 'wrong')]
            preserved = [r for r in produced if (r['checks'].get('preservation') or {}).get('pass')]
            kinds = Counter((r['checks'].get('preservation') or {}).get('kind') for r in produced
                            if (r['checks'].get('preservation') or {}).get('pass') is False)
            # A record from before the stamp filter carries no strict reading: its pass was strict.
            strict = [r for r in produced if ((r['checks'].get('preservation') or {}).get('strict')
                                              or r['checks'].get('preservation') or {}).get('pass')]
            grades = [((r['checks'].get('preservation') or {}).get('grade') or {}) for r in produced]
            graded = Counter(g.get('verdict') for g in grades if g.get('verdict'))
            judged = [r for r in produced if (r['checks'].get('opens') or {}).get('pass') is not None]
            times = [r['seconds'] for r in rows if r['commandCount']]
            table[tool] = dict(
                tasks=len(rows), **{name: counts[name] for name in CORPUS_RESULTS},
                attempted=attempted, passRate=rate(counts['pass'], attempted),
                wrongRate=rate(counts['wrong'], attempted), passRateOfTasks=rate(counts['pass'], len(rows)),
                produced=len(produced), preserved=len(preserved),
                preservationRate=rate(len(preserved), len(produced)),
                preservationFailuresByKind={name: count for name, count in sorted(kinds.items()) if name},
                strictPreserved=len(strict),
                graded=sum(graded.values()), gradePreserved=graded['preserved'],
                gradeRewritten=graded['rewritten'], gradeChanged=graded['changed'],
                stamped=sum(1 for g in grades if (g.get('counts') or {}).get('stamps')),
                passedAndPreserved=sum(1 for r in preserved if r['result'] == 'pass'),
                opensChecked=len(judged), opensFailed=sum(1 for r in judged if not r['checks']['opens']['pass']),
                medianSeconds=round(statistics.median(times), 3) if times else None,
                p95Seconds=round(percentile95(times), 3) if times else None,
                commands=sum(r['commandCount'] for r in rows),
                stdoutBytes=sum(r['stdoutBytes'] for r in rows))
        summary[scope] = table
    return summary


def percent(value):
    return 'n/a' if value is None else f'{value * 100:.1f}%'


def corpus_table(table):
    """One row per tool. Preserved is the stamp-aware pass; the graded columns say what the
    produced files carry outside their write set: nothing (preserved), only same-content
    rewrites (rewritten only) or at least one difference a reader could notice (changed)."""
    rows = ['| Tool | Tasks | Pass | Wrong | Declined | Unsupported | Error | Pass rate | Wrong rate | '
            'Preserved | Preservation rate | Foreign member | Edited part | Graded preserved | Rewritten only | '
            'Changed | Stamped | Median s | p95 s |',
            '| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | '
            '---: | ---: | ---: | ---: |']
    for tool, value in table.items():
        median = 'n/a' if value['medianSeconds'] is None else f'{value["medianSeconds"]:.3f}'
        p95 = 'n/a' if value['p95Seconds'] is None else f'{value["p95Seconds"]:.3f}'
        rows.append(f'| {tool} | {value["tasks"]} | {value["pass"]} | {value["wrong"]} | {value["declined"]} | '
                    f'{value["unsupported"]} | {value["error"]} | {percent(value["passRate"])} | '
                    f'{percent(value["wrongRate"])} | {value["preserved"]}/{value["produced"]} | '
                    f'{percent(value["preservationRate"])} | '
                    f'{value["preservationFailuresByKind"].get("foreign_member", 0)} | '
                    f'{value["preservationFailuresByKind"].get("edited_part", 0)} | '
                    f'{value.get("gradePreserved", 0)} | {value.get("gradeRewritten", 0)} | '
                    f'{value.get("gradeChanged", 0)} | {value.get("stamped", 0)} | {median} | {p95} |')
    return '\n'.join(rows)


CORPUS_CAVEATS = (
    'The tasks are authored by this runner from each file\'s own bytes, not by a customer and not by '
    'any tool under test. Every tool receives the same task statement and the same source copy.',
    'Each adapter is our reading of that tool\'s documented route for the task. A tool marked '
    'unsupported may still reach the same end another way, and a better adapter may exist.',
    'A pass means the requested change is present and this runner\'s own ZIP and XML readers found no '
    'other content difference. It is not a statement about visual fidelity, and nothing here was '
    'opened in Word, Excel or PowerPoint.',
    'Preservation is measured against this repository\'s contract: members the change does not need '
    'stay byte-identical, and the edited members stay semantically identical outside the change. '
    'The provenance and save stamps a tool adds are set aside first, and nothing else is: UFO and '
    'OfficeCLI custom properties with the one content-type override and root relationship a new '
    'properties part needs, the modified time, last author and revision of the core properties, and '
    'the application name and statistics of the extended properties. Each record keeps the strict '
    'byte check without that filter. A tool that rewrites a package is not thereby wrong for its own users.',
    'Every produced file is also graded on what it carries outside its write set: preserved when '
    'nothing differs, rewritten only when every difference is a same-content rewrite or bookkeeping '
    'no reader sees (relationship ids, a calculation chain, a new slide\'s own copied parts), changed '
    'when a person or a program reading the file could notice at least one. A workbook is read by '
    'what its cells, formats, rules and layout say, not by how its parts are numbered.',
    'The opens check is this runner\'s own OPC parse and relationship resolution. It is a cheap '
    'sanity check, not native application acceptance, which runs separately over the outputs index.',
    'Latency is wall time for the whole task, including a tool\'s own reads, its retries and the '
    'oracle work, on one sequential cold process per command. It is not per-command startup time '
    'and no model tokens were measured.',
    'PDF files in the manifest are recorded as skipped, not compared: no tool in this set offers a '
    'PDF route, and this runner\'s independent PDF reader follows classic cross-reference tables '
    'with the fixture\'s own font layout, so it cannot judge the corpus PDFs.',
    'One run per task, on one host, with one version of every tool.',
)


def corpus_report(summary, records, manifest, reproduce):
    """The publishable view: per family, per tool, then every wrong result with its reason."""
    lines = [f'# Real-world corpus comparison {manifest["startedUtc"]}', '',
             f'{manifest["files"]} files from `{manifest["manifest"]}` produced {manifest["tasks"]} tasks '
             f'over {len(manifest["families"])} families, run against {len(manifest["tools"])} tools.', '']
    lines += ['## Tools', '', '| Tool | Version | Identity |', '| --- | --- | --- |']
    for tool, identity in manifest['tools'].items():
        lines.append(f'| {tool} | {identity.get("version", "")} | '
                     f'{identity.get("sha256", identity.get("environment", ""))} |')
    for scope, table in summary.items():
        lines += ['', f'## {"Every family" if scope == "all" else scope.upper()}', '', corpus_table(table)]
    wrong = [r for r in records if r['result'] == 'wrong']
    lines += ['', f'## Wrong results ({len(wrong)})', '']
    if wrong:
        lines += ['| Tool | Family | File | Task | Reason |', '| --- | --- | --- | --- | --- |']
        for row in wrong:
            reason = (row['reason'] or '').replace('|', '\\|').replace('\n', ' ')[:300]
            lines.append(f'| {row["tool"]} | {row["family"]} | {row["file"]} | {row["task"]} | {reason} |')
    else:
        lines.append('No tool produced a file whose content this runner judged wrong.')
    lines += ['', '## Files and tasks the plan left out', '',
              'A file appears here when it yields no task at all, and also when it yields one task '
              'but not the other.', '']
    if manifest['skipped']:
        lines += ['| File | Family | Reason |', '| --- | --- | --- |']
        for row in manifest['skipped']:
            lines.append(f'| {row["file"]} | {row["family"]} | {"; ".join(row["reasons"])[:200]} |')
    else:
        lines.append('Every manifest row produced every task its family states.')
    lines += ['', '## What this does not establish', '']
    lines += [f'- {caveat}' for caveat in CORPUS_CAVEATS]
    lines += ['', '## Reproduce', '', '```bash', *reproduce, '```', '']
    return '\n'.join(lines)


def corpus_index(records, manifest):
    """The layout the native Office acceptance step consumes: one row per published file."""
    entries = []
    for record in records:
        if not record['output']:
            continue
        entries.append(dict(path=record['output']['path'], sha256=record['output']['sha256'],
                            bytes=record['output']['bytes'], sourcePath=manifest['sources'][record['file']],
                            sourceSha256=record['sourceSha256'], tool=record['tool'],
                            toolVersion=manifest['tools'][record['tool']].get('version'),
                            file=record['file'], family=record['family'], task=record['task'],
                            status=record['result'],
                            # The cell or the row this task wrote, where the task states one:
                            # the native step never attributes such a cell to the source.
                            sheet=record.get('sheet'), cell=record.get('cell'), row=record.get('row'),
                            preservation=(record['checks'].get('preservation') or {}).get('pass'),
                            opens=(record['checks'].get('opens') or {}).get('pass')))
    return dict(schemaVersion=CORPUS_VERSION, run=manifest['startedUtc'], layout='outputs/<tool>/<file>/<task>.<ext>',
                note='Every published file, whatever this runner judged it. Native application acceptance '
                     'is a separate step over these paths.', entries=entries)


def corpus_main(args):
    started = datetime.now(timezone.utc)
    stamp = started.strftime('%Y%m%dT%H%M%SZ')
    manifest_path = args.corpus.resolve()
    root = args.corpus_root.resolve()
    families = tuple(f.strip() for f in args.corpus_families.split(',') if f.strip()) if args.corpus_families else CORPUS_FAMILIES
    venv = args.libraries.resolve()
    identities = provision(args.tools_dir.resolve(), args.download)
    launcher = args.ufo.absolute()
    if not launcher.is_file():
        raise ValueError(f'UFO app image missing: {launcher}')
    identities = {'ufo': dict(path=str(launcher), sha256=sha(launcher.read_bytes()),
                              artifact=launcher_artifact(launcher), license='proprietary'), **identities}
    identities.update(provision_libraries(venv, args.setup_libraries))
    tools = [name for name in CORPUS_TOOLS if name in identities]
    output = (args.output or REPO / 'build/comparisons' / f'corpus-{stamp}').resolve()
    output.mkdir(parents=True, exist_ok=False)
    outputs = output / 'outputs'
    outputs.mkdir()
    rows, missing = corpus_files(manifest_path, root, families, args.corpus_limit)
    planned, unplanned = corpus_plan(rows)
    skipped = missing + unplanned
    with tempfile.TemporaryDirectory(prefix='setup-', dir=output) as temporary:
        prefix, network = network_mode(Path(temporary))
        for name, identity in identities.items():
            runner = Runner(identity['path'], temporary, prefix, timeout=CORPUS_TIMEOUT)
            if name in LIBRARY_FAMILY:
                identity.setdefault('setupCommands', [])
                continue
            version = runner.run(['--version'])
            if version['returncode']:
                raise ValueError(f'{name} version command failed')
            identity['version'] = version['stdout'].strip()
            identity['setupCommands'] = runner.commands
    manifest = dict(
        schemaVersion=CORPUS_VERSION, startedUtc=started.isoformat().replace('+00:00', 'Z'),
        manifest=str(manifest_path.relative_to(REPO)) if manifest_path.is_relative_to(REPO) else str(manifest_path),
        corpusRoot=str(root), families=list(families), files=len(rows), tasks=len(planned),
        tools={name: identities[name] for name in tools}, skipped=skipped,
        sources={row['id']: row['relative'] for row in rows},
        runnerSha256=sha(Path(__file__).read_bytes()), network=network,
        bounds=dict(maxSourceBytes=CORPUS_MAX_BYTES, commandTimeoutSeconds=CORPUS_TIMEOUT,
                    maxCaptureBytes=MAX_CAPTURE),
        libraryPins={name: version for name, version, _ in LIBRARY_PINS},
        caveats=list(CORPUS_CAVEATS), tasks_planned=planned)
    (output / 'manifest.json').write_bytes(json_bytes(manifest))
    records = []
    with (output / 'records.jsonl').open('w', encoding='utf-8') as stream:
        for position, task in enumerate(planned):
            order = tools[position % len(tools):] + tools[:position % len(tools)]
            data = (root / task['sourcePath']).read_bytes()
            for tool in order:
                record = execute_corpus_task(task, tool, identities[tool], data, prefix, outputs, CORPUS_TIMEOUT)
                records.append(record)
                stream.write(json.dumps(record, sort_keys=True, ensure_ascii=False) + '\n')
                stream.flush()
            print(f'[{position + 1}/{len(planned)}] {task["id"]}: '
                  + ' '.join(f'{r["tool"]}={r["result"]}' for r in records[-len(order):]),
                  file=sys.stderr, flush=True)
    summary = corpus_summarize(records, tools)
    reproduce = [
        f'# once, with network, to create the pinned library environment under build/',
        f'python3 tools/corpus/compare_agent_tools.py --corpus {manifest["manifest"]} --setup-libraries',
        '# the measured run',
        f'python3 tools/corpus/compare_agent_tools.py --corpus {manifest["manifest"]} \\',
        f'  --corpus-root {manifest["corpusRoot"]} \\',
        f"  --ufo '{args.ufo}'",
    ]
    (output / 'summary.json').write_bytes(json_bytes(dict(
        schemaVersion=CORPUS_VERSION, manifest=manifest, summary=summary,
        wrong=[dict(tool=r['tool'], family=r['family'], file=r['file'], task=r['task'], reason=r['reason'],
                    checks=r['checks'], output=r['output']) for r in records if r['result'] == 'wrong'],
        reproduce=reproduce)))
    (output / 'summary.md').write_text(corpus_report(summary, records, manifest, reproduce), encoding='utf-8')
    (outputs / 'index.json').write_bytes(json_bytes(corpus_index(records, manifest)))
    print(corpus_table(summary['all']))
    print(f'\nReceipts: {output}')
    return 0


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--download',action='store_true',help='download pinned release binaries if absent; all runs after provisioning are offline-capable')
    parser.add_argument('--ufo',type=Path,default=APP)
    parser.add_argument('--ufo-brief',action='store_true',help='use UFO --brief display projection; report alongside a default-output run, without claiming all adapters are shortest')
    parser.add_argument('--ufo-compact',action='store_true',help='use UFO --compact display projection, the smallest form; needs a launcher built from source that carries the flag')
    parser.add_argument('--tools-dir',type=Path,default=REPO/'build/compare-tools')
    parser.add_argument('--output',type=Path)
    parser.add_argument('--corpus',type=Path,help='run the real-world corpus mode over this manifest instead of the pinned 24-task manifest')
    parser.add_argument('--corpus-root',type=Path,default=REPO/'local-corpus/realworld',help='directory the manifest local_path column is relative to')
    parser.add_argument('--corpus-families',default=','.join(CORPUS_FAMILIES))
    parser.add_argument('--corpus-limit',type=int,help='stop after this many manifest files; for a smoke run only')
    parser.add_argument('--libraries',type=Path,default=REPO/'build/compare-libraries/venv',help='private virtual environment holding the pinned python-docx, openpyxl and python-pptx')
    parser.add_argument('--setup-libraries',action='store_true',help='create the pinned library environment; needs network once, like --download')
    args = parser.parse_args(argv)
    try:
        if args.ufo_brief and args.ufo_compact: raise ValueError('choose --ufo-brief or --ufo-compact, not both')
        if args.corpus: return corpus_main(args)
        ufo_display = 'compact' if args.ufo_compact else ('brief' if args.ufo_brief else None)
        identities = provision(args.tools_dir.resolve(),args.download)
        launcher = args.ufo.absolute()
        if not launcher.is_file(): raise ValueError(f'UFO app image missing: {launcher}')
        identities = {'ufo':dict(path=str(launcher),sha256=sha(launcher.read_bytes()),artifact=launcher_artifact(launcher)),**identities}
        output = (args.output or REPO/'build/compare'/datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S.%fZ')).resolve()
        output.mkdir(parents=True,exist_ok=False)
        inputs = fixtures()
        manifest = dict(version=MANIFEST_VERSION,ufo_brief=args.ufo_brief,ufo_display=ufo_display,tasks=TASKS,fixtures={k:sha(v) for k,v in inputs.items()},
                        fixture_builder_sha256=sha((REPO/'tools/cli/evaluate_sample.py').read_bytes()),
                        preservation='Untouched package members byte-identical once the provenance and save stamps a tool adds are set aside (and nothing else); unrequested content in edited parts semantically unchanged.',
                        timing='One sequential cold process per command; includes adapter reads, failed attempts and oracle work; excludes provisioning/discovery; latency excludes unsupported zero-command tasks.',
                        authorship='Vendor-authored synthetic tasks, four lineages, no holdout, no model, no model tokens, no native Office or visual acceptance.',
                        adapter_notes={
                            'officecli':'Documented synchronous flush: OFFICECLI_RESIDENT_FLUSH=each. Writes use a staging copy. DOCX replacement tries scoped whole-text find/replace, then scoped changed-span find/replace. PPTX tries shape text, then run text. Every attempt counts.',
                            'docx-cli':'Targeted Markdown paragraph slice, outline JSON. Replacement tries scoped whole-text replace, then scoped changed-span replace. -o writes a copy. Every attempt counts.',
                            'ufo':'Structured reads supply hashes and native coordinates before edits. Sheet rename uses edit sheets --rename, which the September 15 launcher lacked. The PDF page command has no hash-precondition option; the adapter first reads the page inventory.',
                            'genoffice':'The genoffice command line from the pinned Debian package, run in place on its bundled Electron runtime with its audit log off and its path policy confined to the task folder. Writes use docs, sheet and slides apply with --out. DOCX and PPTX replacement and the tracked replacement try the whole text, then the changed characters, because its find does not match across differently formatted runs; the table cell tries a find scoped to the table, then rewrites the table block from its own HTML read. Cell writes carry the current value as expectedValue. PDF text is read by converting to Word locally. Every attempt counts.',
                        },
                        documentation={name:[f'https://github.com/{p["repository"]}/blob/{p["tag"]}/README.md',p['url']] for name,p in PINS.items()})
        (output/'manifest.json').write_bytes(json_bytes(manifest))
        (output/'inputs').mkdir()
        for family,data in inputs.items(): (output/'inputs'/('input.'+family)).write_bytes(data)
        with tempfile.TemporaryDirectory(prefix='setup-',dir=output) as tmp:
            prefix,network = network_mode(Path(tmp))
            for name,identity in identities.items():
                r = Runner(identity['path'],tmp,prefix)
                version = r.run(['--version'])
                if version['returncode']: raise ValueError(f'{name} version command failed')
                identity['version'] = version['stdout'].strip()
                r.run(['--help'])
                identity['setup_commands'] = r.commands
        records = []
        # Rotate tools per task to reduce persistent first/last-order effects.
        names = list(identities)
        for i,task in enumerate(TASKS):
            for name in names[i%len(names):]+names[:i%len(names)]:
                record = execute_task(task,name,identities[name]['path'],inputs[task['family']],prefix,output,ufo_display=ufo_display)
                records.append(record)
                print(f'{task["id"]} {name}: {record["result"]}',file=sys.stderr,flush=True)
        summary = dict(schema_version=1,manifest_sha256=sha(json_bytes(manifest)),runner_sha256=sha(Path(__file__).read_bytes()),
                       network=network,limitations=manifest['authorship'],timing=manifest['timing'],tools=summarize(records,identities))
        (output/'summary.json').write_bytes(json_bytes(summary))
        print(markdown_table(summary['tools']))
        print(f'\nReceipts: {output}')
        return 1 if any(r['result'] == 'error' for r in records) else 0
    except (ValueError,OSError,subprocess.SubprocessError) as exc:
        print(f'comparison: {exc}',file=sys.stderr); return 2


if __name__ == '__main__':
    raise SystemExit(main())
